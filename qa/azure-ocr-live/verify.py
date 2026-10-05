"""Explicit live Azure OCR check using only generated synthetic documents."""

from __future__ import annotations

import hashlib
import json
import os
import re
import socket
import subprocess
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path

import httpx
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from docling_azure_ocr.config import OcrError, api_key, load_profile  # noqa: E402

import settings  # noqa: E402,F401 -- Load .env before selecting credentials.
from docling_desk.documents.ocr_profiles import make_runtime  # noqa: E402

OUTPUT = ROOT / "qa/azure-ocr-live" / datetime.now().strftime("%Y%m%d-%H%M%S")
OUTPUT.mkdir(parents=True)
REPORT = {"verified_at": datetime.now().isoformat(), "provider": "azure_read", "cases": []}
EXPECTED = ["画像OCR", "在庫", "73921", "120", "2026"]


def compact(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value))


def check_text(text):
    normalized = compact(text)
    matches = {word: compact(word) in normalized for word in EXPECTED}
    if not all(matches.values()):
        raise AssertionError(
            "Expected synthetic text missing: " + json.dumps(matches, ensure_ascii=False)
        )
    return matches


def save_report():
    (OUTPUT / "verification.json").write_text(json.dumps(REPORT, ensure_ascii=False, indent=2))


def documents(image_path):
    import pymupdf
    from docx import Document
    from docx.shared import Inches as WordInches
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage
    from pptx import Presentation
    from pptx.util import Inches

    paths = []
    path = OUTPUT / "synthetic-azure-image.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), Inches(0.5), Inches(1), width=Inches(9))
    deck.save(path)
    paths.append(("pptx", path))
    for table in (False, True):
        label = "docx-table" if table else "docx-body"
        path = OUTPUT / f"synthetic-azure-{label}.docx"
        document = Document()
        target = document.add_table(rows=1, cols=1).cell(0, 0) if table else document
        target.add_paragraph().add_run().add_picture(str(image_path), width=WordInches(6))
        document.save(path)
        paths.append((label, path))
    path = OUTPUT / "synthetic-azure-image.xlsx"
    workbook = Workbook()
    workbook.active.title = "OCR Test"
    workbook.active.add_image(SheetImage(str(image_path)), "B2")
    workbook.save(path)
    paths.append(("xlsx", path))
    path = OUTPUT / "synthetic-azure-image.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page(width=750, height=400)
    page.insert_image(pymupdf.Rect(25, 40, 725, 300), filename=str(image_path))
    pdf.save(path)
    pdf.close()
    paths.append(("pdf-image-only", path))
    return paths


def main():
    profile = load_profile("azure_read")
    profile.check_credentials(api_key())
    assert profile.auth == "api_key", "The configured API key is required for this live check."
    image_path = OUTPUT / "synthetic-ocr.png"
    bitmap = Image.new("RGB", (1500, 520), "white")
    font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Unicode.ttf", 52)
    drawer = ImageDraw.Draw(bitmap)
    for i, line in enumerate(
        [
            "画像OCRの動作確認",
            "商品番号：73921",
            "在庫数：120個",
            "次回入荷：2026年10月15日",
            "IMAGE OCR TEST 73921",
        ]
    ):
        drawer.text((45, 25 + i * 95), line, font=font, fill="black")
    bitmap.save(image_path)
    folder = OUTPUT / "direct-image"
    folder.mkdir()
    (folder / "original.png").write_bytes(image_path.read_bytes())
    runtime = make_runtime(
        profile,
        folder,
        hashlib.sha256(image_path.read_bytes()).hexdigest(),
        "synthetic-live-direct",
    )
    try:
        result = runtime.read(bitmap, "image:1", {"coordinate_frame": "embedded_image_pixels"})
        text = result.get("text") or " ".join(word["text"] for word in result["words"])
        check_text(text)
        records_before = [
            p.read_bytes() for p in sorted((folder / "ocr-private/operations").glob("*.json"))
        ]
        runtime.read(bitmap, "image:1", {"coordinate_frame": "embedded_image_pixels"})
        records_after = [
            p.read_bytes() for p in sorted((folder / "ocr-private/operations").glob("*.json"))
        ]
        assert records_after == records_before, "Cache reuse changed the submission ledger."
        REPORT["cases"].append(
            {
                "case": "png-direct",
                "status": "passed",
                "recognized_text": text,
                "cache_reused_without_submission": True,
            }
        )
        save_report()
        print("PASS: actual Azure image OCR and cache reuse", flush=True)
    finally:
        runtime.transport.close()

    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    environment = os.environ.copy()
    environment["DOCLING_DATA_DIR"] = str(OUTPUT / "jobs")
    with (OUTPUT / "server.log").open("w") as log:
        server = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "uvicorn",
                "docling_desk.app:app",
                "--host",
                "127.0.0.1",
                "--port",
                str(port),
                "--no-access-log",
                "--env-file",
                str(ROOT / ".env"),
            ],
            cwd=ROOT,
            env=environment,
            stdout=log,
            stderr=log,
        )
        try:
            with httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=30) as client:
                deadline = time.monotonic() + 40
                while time.monotonic() < deadline:
                    assert server.poll() is None, "Verification server exited."
                    try:
                        if client.get("/health/live").status_code == 200:
                            break
                    except httpx.HTTPError:
                        pass
                    time.sleep(0.3)
                else:
                    raise AssertionError("Verification server did not start.")
                capabilities = client.get("/api/ocr").json()
                assert capabilities == {"azure_available": True, "default_provider": "local"}
                REPORT["ocr_api"] = capabilities
                for label, path in documents(image_path):
                    with path.open("rb") as source:
                        response = client.post(
                            "/api/upload",
                            data={"ocr_provider": "azure_read"},
                            files={"file": (path.name, source)},
                        )
                    assert response.status_code == 202, (
                        f"{label}: upload status {response.status_code}"
                    )
                    job_id = response.json()["id"]
                    deadline = time.monotonic() + 240
                    while time.monotonic() < deadline:
                        job = client.get(f"/api/jobs/{job_id}").json()
                        if job["state"] not in ("queued", "running"):
                            break
                        time.sleep(0.8)
                    assert job["state"] in ("success", "partial"), (
                        f"{label}: {job['state']}, {job.get('error_code')}"
                    )
                    job_folder = OUTPUT / "jobs" / job_id
                    text = (job_folder / "text.txt").read_text()
                    matches = check_text(text)
                    for name in ("text.txt", "document.md", "extracted.html"):
                        exported = (job_folder / name).read_text()
                        assert "73921" in exported
                        assert "response_sha256" not in exported
                        assert "cognitiveservices.azure.com" not in exported
                    for name in ("rag.jsonl", "rag-index.jsonl", "rag-docling.jsonl"):
                        records = [
                            json.loads(line)
                            for line in (job_folder / name).read_text().splitlines()
                            if line
                        ]
                        assert records and all("response_sha256" not in r["text"] for r in records)
                    chunks = [
                        json.loads(line)
                        for line in (job_folder / "rag-index.jsonl").read_text().splitlines()
                        if line
                    ]
                    check_text("\n".join(chunk["text"] for chunk in chunks))
                    evidence = [item for chunk in chunks for item in chunk.get("ocr_evidence", [])]
                    assert evidence, f"{label}: OCR evidence missing from RAG"
                    digest = hashlib.sha256(path.read_bytes()).hexdigest()
                    assert all(
                        e["source_sha256"] == digest and e["provider"] == "azure_read"
                        for e in evidence
                    )
                    if label != "pdf-image-only":
                        assert any(
                            e.get("picture_ref", "").startswith("#/pictures/") for e in evidence
                        )
                    if label in ("pptx", "xlsx", "pdf-image-only"):
                        assert any(1 in chunk["pages"] for chunk in chunks)
                    for relative in ("text.txt", "rag-index.jsonl", "document.json"):
                        assert client.get(f"/files/{job_id}/{relative}").status_code == 200
                    operations = [
                        json.loads(p.read_text())
                        for p in (job_folder / "ocr-private/operations").glob("*.json")
                    ]
                    assert operations and all(op["state"] == "completed" for op in operations)
                    case = {
                        "case": label,
                        "status": "passed",
                        "job_state": job["state"],
                        "job_id": job_id,
                        "recognized_text": text,
                        "expected_tokens": matches,
                        "rag_chunks": len(chunks),
                        "ocr_evidence": len(evidence),
                        "azure_operations": len(operations),
                        "duration_seconds": job["duration"],
                        "clean_text_and_exports": True,
                    }
                    REPORT["cases"].append(case)
                    save_report()
                    print(
                        f"PASS: {label}, image text in RAG, {len(evidence)} OCR evidence records, {job['duration']}s",
                        flush=True,
                    )
        finally:
            server.terminate()
            try:
                server.wait(timeout=8)
            except subprocess.TimeoutExpired:
                server.kill()
                server.wait(timeout=5)
    REPORT["status"] = "passed"
    save_report()
    key = api_key().get_secret_value()
    assert all(
        key not in p.read_text(errors="ignore")
        for p in OUTPUT.rglob("*")
        if p.is_file() and p.suffix in (".json", ".jsonl", ".log", ".txt")
    )
    print("PASS: API key absent from verification artifacts", flush=True)
    print("Report: " + str(OUTPUT / "verification.json"), flush=True)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        REPORT["status"] = "failed"
        REPORT["error_type"] = type(exc).__name__
        REPORT["error_code"] = exc.code if isinstance(exc, OcrError) else "verification_failed"
        save_report()
        print("FAILED: " + REPORT["error_type"] + " / " + REPORT["error_code"], flush=True)
        print("Report: " + str(OUTPUT / "verification.json"), flush=True)
        sys.exit(1)
