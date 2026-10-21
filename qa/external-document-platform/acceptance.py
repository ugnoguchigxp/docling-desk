"""Real HTTP, MCP transport, Docling processor, browser and process recovery.

Only fresh synthetic data and private ephemeral credentials are used. The model
provider is disabled. Run from the repository root after frontend build/export.
"""

from __future__ import annotations

import asyncio
import base64
import hashlib
import hmac
import json
import os
import shutil
import socket
import sqlite3
import subprocess
import tempfile
import time
from datetime import datetime
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo

import httpx
from docx import Document
from mcp import ClientSession
from mcp.client.streamable_http import streamablehttp_client
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / "qa/external-document-platform"


def port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def signed(client, actor="writer"):
    def encode(value):
        return (
            base64.urlsafe_b64encode(json.dumps(value, separators=(",", ":")).encode())
            .decode()
            .rstrip("=")
        )

    now = int(time.time())
    header = encode({"alg": "HS256", "typ": "JWT", "kid": "v1"})
    payload = encode(
        {
            "iss": client["issuer"],
            "aud": "docling-desk-api",
            "client_id": client["id"],
            "sub": actor,
            "iat": now,
            "exp": now + 300,
            "scopes": client["scopes"],
        }
    )
    signature = (
        base64.urlsafe_b64encode(
            hmac.digest(client["keys"]["v1"].encode(), f"{header}.{payload}".encode(), "sha256")
        )
        .decode()
        .rstrip("=")
    )
    return {
        "Authorization": "Bearer " + client["token"],
        "X-Knowledge-Actor": f"{header}.{payload}.{signature}",
    }


def wait_ready(url, process, seconds=60):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise RuntimeError("process exited: " + url)
        try:
            if httpx.get(url, timeout=2).is_success:
                return
        except httpx.HTTPError:
            pass
        time.sleep(0.25)
    raise RuntimeError("readiness timeout: " + url)


async def scenario(client_config, config, api_url, api_env, processes, logs, temporary):
    results = {}
    write_client = client_config[1]
    client = httpx.Client(base_url=api_url, timeout=40)

    def api(path, method="GET", body=None, headers=None):
        return client.request(
            method, path, json=body, headers={**signed(write_client), **(headers or {})}
        )

    async def call(session, name, arguments=None):
        result = await session.call_tool(name, arguments or {})
        if result.isError:
            raise RuntimeError(str(result.content))
        if result.structuredContent:
            return result.structuredContent
        return json.loads(result.content[0].text)

    async def ready_job(job_id):
        for _ in range(180):
            data = api("/api/v1/jobs/" + job_id).json()
            if data["state"] == "completed":
                return
            if data["state"] == "failed":
                raise RuntimeError("job failed: " + str(data["error_code"]))
            await asyncio.sleep(0.5)
        raise RuntimeError("job timeout")

    # Actual PDF, Office and Markdown registration through the existing multipart API.
    pdf_path = temporary / "synthetic.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=A4)
    pdf.drawString(50, A4[1] - 80, "Synthetic inventory evidence 120 items.")
    pdf.showPage()
    pdf.save()
    office = Document()
    office.add_heading("Synthetic inventory", 1)
    office.add_paragraph("Synthetic inventory evidence 120 items.")
    office_path = temporary / "synthetic.docx"
    office.save(office_path)
    markdown_path = temporary / "evidence.md"
    markdown_path.write_text("# Synthetic evidence\n\nSynthetic inventory evidence 120 items.\n")
    references = []
    for path in (pdf_path, office_path, markdown_path):
        response = client.post(
            "/api/v1/sources",
            headers={**signed(write_client), "Idempotency-Key": str(uuid4())},
            files={"file": (path.name, path.read_bytes())},
            data={"metadata": json.dumps({"title": path.name, "collection_id": "synthetic"})},
        )
        response.raise_for_status()
        data = response.json()
        await ready_job(data["job_id"])
        references.append(data)
    results["real_docling_formats"] = ["pdf", "docx", "markdown"]
    private = api(
        "/api/v1/markdown-sources",
        "POST",
        {
            "body": "Private scope evidence",
            "metadata": {"title": "Private", "collection_id": "other"},
        },
        {"Idempotency-Key": str(uuid4())},
    )
    private.raise_for_status()
    private_id = private.json()["source_id"]

    print("Registration and real Docling extraction passed", flush=True)

    async with streamablehttp_client(
        config["public_url"] + "/mcp/",
        headers={"Authorization": "Bearer " + config["users"]["writer"]["mcp_token"]},
    ) as (read, write, _):
        async with ClientSession(read, write) as mcp:
            await mcp.initialize()
            tools = await mcp.list_tools()
            assert "commit_document_draft" in {item.name for item in tools.tools}
            search = await call(
                mcp,
                "search_documents",
                {
                    "query": "inventory",
                    "scope": {"collection_ids": ["synthetic"]},
                    "mode": "hybrid",
                },
            )
            assert search["results"] and search["effective_mode"] == "text"
            hit = search["results"][0]
            await call(
                mcp,
                "get_document_context",
                {
                    "retrieval_id": search["retrieval_id"],
                    "references": [
                        {
                            key: hit[key]
                            for key in (
                                "source_id",
                                "source_revision",
                                "evidence_revision",
                                "context_id",
                            )
                        }
                    ],
                },
            )
            raw = "\ufeff---\r\ntitle: raw title\r\n---\r\n\r\n# Blog draft\r\n\r\nOriginal paragraph with inventory evidence.\r\n\r\n```mermaid\r\ngraph TD; A-->B\r\n```\r\n\r\n$$x^2$$\r\n\r\n:::message\r\nCustom block\r\n:::\r\n\r\n<Component value={neverExecute()} />\r\n"
            create_key = str(uuid4())
            draft = await call(
                mcp,
                "create_document_draft",
                {
                    "arguments": {
                        "body": raw,
                        "metadata": {
                            "title": "Synthetic blog",
                            "collection_id": "synthetic",
                            "source_kind": "wiki",
                        },
                        "idempotency_key": create_key,
                    }
                },
            )
            replay = await call(
                mcp,
                "create_document_draft",
                {
                    "arguments": {
                        "body": raw,
                        "metadata": {
                            "title": "Synthetic blog",
                            "collection_id": "synthetic",
                            "source_kind": "wiki",
                        },
                        "idempotency_key": create_key,
                    }
                },
            )
            assert replay["draft_id"] == draft["draft_id"]
            current = api("/api/v1/drafts/" + draft["draft_id"]).json()
            assert current["body"] == raw
            draft = await call(
                mcp,
                "edit_document_draft",
                {
                    "arguments": {
                        "draft_id": draft["draft_id"],
                        "etag": draft["etag"],
                        "idempotency_key": str(uuid4()),
                        "base_sha256": draft["sha256"],
                        "edits": [
                            {
                                "expected": "inventory evidence",
                                "replacement": "verified inventory evidence",
                            }
                        ],
                    }
                },
            )
            image_bytes = base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl6kVcAAAAASUVORK5CYII="
            )
            attachment = await call(
                mcp,
                "add_document_attachment",
                {
                    "arguments": {
                        "draft_id": draft["draft_id"],
                        "etag": draft["etag"],
                        "filename": "receipt.png",
                        "media_type": "image/png",
                        "body_base64": base64.b64encode(image_bytes).decode(),
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            draft = await call(
                mcp,
                "update_document_draft",
                {
                    "arguments": {
                        "draft_id": draft["draft_id"],
                        "etag": attachment["etag"],
                        "body": draft["body"]
                        + "\r\n![receipt]("
                        + attachment["reference"]
                        + ")\r\n",
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            retrieved = await call(
                mcp,
                "get_draft_attachment",
                {
                    "arguments": {
                        "draft_id": draft["draft_id"],
                        "attachment_id": attachment["attachment_id"],
                    }
                },
            )
            assert retrieved["sha256"] == hashlib.sha256(image_bytes).hexdigest()
            results["http_mcp_same_draft"] = True
            denied = await mcp.call_tool("get_document", {"arguments": {"source_id": private_id}})
            assert denied.isError
            listing = await call(mcp, "list_documents", {"arguments": {}})
            assert private_id not in {document["source_id"] for document in listing["documents"]}
            results["out_of_scope_mcp_denied"] = True

            forged = await mcp.call_tool(
                "update_document_draft",
                {
                    "arguments": {
                        "draft_id": draft["draft_id"],
                        "body": "forged",
                        "user_id": "admin",
                        "scope": {"collection_id": "other"},
                        "etag": draft["etag"],
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            assert forged.isError
            results["forged_mcp_identity_denied"] = True

            # UI and host pairing operate with credentials kept on the host, never the browser.
            browser_input = {
                "public_url": config["public_url"],
                "connector_token": config["connector_token"],
                "draft_id": draft["draft_id"],
                "screenshot": str(QA / "editor.png"),
                "result_file": str(temporary / "browser-result.json"),
            }
            browser_file = temporary / "browser.json"
            browser_file.write_text(json.dumps(browser_input))
            browser_file.chmod(0o600)
            completed = await asyncio.to_thread(
                subprocess.run,
                ["node", str(QA / "browser.mjs"), str(browser_file)],
                cwd=ROOT,
                capture_output=True,
                text=True,
                timeout=90,
            )
            if completed.returncode:
                raise RuntimeError("browser failed: " + completed.stderr)
            results.update(json.loads((temporary / "browser-result.json").read_text()))
            latest = api("/api/v1/drafts/" + draft["draft_id"]).json()
            assert latest["state"] == "committed" and "Human receipt" in latest["body"]
            stale = await mcp.call_tool(
                "update_document_draft",
                {
                    "arguments": {
                        "draft_id": draft["draft_id"],
                        "body": "overwrite",
                        "etag": draft["etag"],
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            assert stale.isError
            results["human_ai_conflict"] = True
            source = api("/api/v1/sources/" + latest["source_id"] + "/body").json()
            assert (
                source["body"].startswith("\ufeff---\r\n")
                and "```mermaid\r\ngraph TD" in source["body"]
                and ":::message\r\n" in source["body"]
            )
            assert source["sha256"] == hashlib.sha256(source["body"].encode()).hexdigest()
            media = await call(
                mcp,
                "get_document_attachment",
                {
                    "arguments": {
                        "source_id": source["source_id"],
                        "attachment_id": attachment["attachment_id"],
                    }
                },
            )
            assert base64.b64decode(media["body_base64"]) == image_bytes
            assert attachment["reference"] in source["body"]
            results["attachment_round_trip"] = True

            # Wait until the committed source has evidence and FTS publication.
            for _ in range(120):
                source = api("/api/v1/sources/" + latest["source_id"] + "/body").json()
                if source["status"]["fts_ready"]:
                    break
                if source["status"]["error_code"]:
                    raise RuntimeError("Markdown index failed: " + source["status"]["error_code"])
                await asyncio.sleep(0.5)
            assert source["status"]["fts_ready"]
            search = await call(
                mcp,
                "search_documents",
                {
                    "query": "Human receipt",
                    "scope": {"collection_ids": ["synthetic"]},
                    "mode": "text",
                },
            )
            hit = next(hit for hit in search["results"] if hit["source_id"] == source["source_id"])
            view_args = {
                key: hit[key] for key in ("source_id", "source_revision", "evidence_revision")
            }
            if hit["locator"].get("original_range"):
                view_args["location"] = {
                    "kind": "document",
                    "number": 1,
                    "original_range": hit["locator"]["original_range"],
                }
            view = await call(mcp, "request_document_view", view_args)
            challenge = httpx.post(
                config["public_url"] + "/viewer/challenges", json=view["reference"]
            ).json()
            await call(mcp, "connect_document_view", {"code": challenge["code"]})
            token = httpx.get(
                config["public_url"] + "/viewer/challenges/" + challenge["challenge"]
            ).json()["session"]
            asset = httpx.get(config["public_url"] + "/viewer/session/" + token + "/view/document")
            assert (
                asset.is_success
                and "Human receipt" in asset.text
                and "&lt;Component" in asset.text
                and "<Component" not in asset.text
            )
            image = httpx.get(
                config["public_url"]
                + "/viewer/session/"
                + token
                + "/attachment/"
                + attachment["attachment_id"]
            )
            assert image.is_success and image.content == image_bytes
            assert image.headers["content-disposition"] == "attachment"
            assert "/attachment/" + attachment["attachment_id"] in asset.text
            results["viewer_attachment_round_trip"] = True
            results["search_and_authenticated_viewer"] = True
            results["original_range_available"] = "location" in view_args
            print("MCP, rich editor, preserved raw source and viewer passed", flush=True)
            # Pause heavy processing, persist a fresh draft and committed source, then restart API.
            control_headers = {"Authorization": "Bearer " + api_env["KNOWLEDGE_CONTROL_TOKEN"]}
            paused = client.put(
                "/internal/v1/processing", headers=control_headers, json={"enabled": False}
            )
            paused.raise_for_status()
            recovery = await call(
                mcp,
                "create_document_draft",
                {
                    "arguments": {
                        "body": "Recovery evidence",
                        "metadata": {"title": "Recovery", "collection_id": "synthetic"},
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            commit_key = str(uuid4())
            committed = await call(
                mcp,
                "commit_document_draft",
                {
                    "arguments": {
                        "draft_id": recovery["draft_id"],
                        "etag": recovery["etag"],
                        "source_etag": None,
                        "idempotency_key": commit_key,
                    }
                },
            )
            assert committed["status"]["search"] == "pending"
            processes[1].terminate()
            await asyncio.to_thread(processes[1].wait, 20)
            processes[1] = subprocess.Popen(
                ["bun", str(ROOT / "knowledge-api/src/server.ts")],
                cwd=ROOT,
                env=api_env,
                stdout=logs[1],
                stderr=logs[1],
            )
            await asyncio.to_thread(wait_ready, api_url + "/health/ready", processes[1])
            replay = await call(
                mcp,
                "commit_document_draft",
                {
                    "arguments": {
                        "draft_id": recovery["draft_id"],
                        "etag": recovery["etag"],
                        "source_etag": None,
                        "idempotency_key": commit_key,
                    }
                },
            )
            assert replay["source_revision"] == committed["source_revision"]
            lookup = await call(
                mcp,
                "get_document_operation",
                {
                    "arguments": {
                        "operation": "commit",
                        "target_id": recovery["draft_id"],
                        "key": commit_key,
                    }
                },
            )
            assert lookup["result"]["source_revision"] == committed["source_revision"]
            assert not client.get("/internal/v1/processing", headers=control_headers).json()[
                "enabled"
            ]
            client.put(
                "/internal/v1/processing", headers=control_headers, json={"enabled": True}
            ).raise_for_status()
            await ready_job(committed["job_id"])
            with sqlite3.connect(
                f"file:{temporary / 'state/knowledge.sqlite'}?mode=ro", uri=True
            ) as db:
                assert (
                    db.execute(
                        "SELECT count(*) FROM revisions WHERE source_id=?",
                        (committed["source_id"],),
                    ).fetchone()[0]
                    == 1
                )
            results["restart_replay_and_processing_resume"] = True
            # An obsolete citation cannot silently resolve to the new revision.
            new_draft = await call(
                mcp,
                "create_document_draft",
                {
                    "arguments": {
                        "source_id": source["source_id"],
                        "source_revision": source["source_revision"],
                        "source_etag": source["etag"],
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            await call(
                mcp,
                "commit_document_draft",
                {
                    "arguments": {
                        "draft_id": new_draft["draft_id"],
                        "etag": new_draft["etag"],
                        "source_etag": new_draft["base_etag"],
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            old = await mcp.call_tool("request_document_view", view_args)
            assert old.isError
            assert (
                httpx.get(
                    config["public_url"] + "/viewer/session/" + token + "/manifest"
                ).status_code
                == 409
            )
            results["old_citation_revoked"] = True
    async with streamablehttp_client(
        config["public_url"] + "/mcp/",
        headers={"Authorization": "Bearer " + config["users"]["reader"]["mcp_token"]},
    ) as (read, write, _):
        async with ClientSession(read, write) as mcp:
            await mcp.initialize()
            denied = await mcp.call_tool(
                "create_document_draft",
                {
                    "arguments": {
                        "body": "forged",
                        "metadata": {"title": "x", "collection_id": "synthetic"},
                        "idempotency_key": str(uuid4()),
                    }
                },
            )
            assert denied.isError
            results["read_only_mcp_denied"] = True
    client.close()
    return results


def main():
    QA.mkdir(exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix="docling-platform-"))
    processes = []
    logs = []
    try:
        api_port, processor_port, gateway_port = port(), port(), port()
        processor_token, control_token = str(uuid4()) + str(uuid4()), str(uuid4()) + str(uuid4())
        clients = []
        for name, actions in (("reader", ["read"]), ("editor", ["read", "write"])):
            clients.append(
                {
                    "id": name,
                    "issuer": name,
                    "mode": "actor",
                    "token": str(uuid4()) + str(uuid4()),
                    "keys": {"v1": str(uuid4()) + str(uuid4())},
                    "actions": actions,
                    "scopes": [{"collection_id": "synthetic"}, {"collection_id": "other"}],
                }
            )
        state = temporary / "state"
        state.mkdir()
        (state / "clients.json").write_text(json.dumps(clients))
        (state / "clients.json").chmod(0o600)
        artifacts = temporary / "artifacts"
        artifacts.mkdir()
        common = {
            **os.environ,
            "PYTHONPATH": str(ROOT / "src"),
            "DOCLING_ENV_FILE": "/dev/null",
            "DOCLING_AZURE_OCR_ENDPOINT": "",
            "AZURE_DOCUMENT_INTELLIGENCE_API_KEY": "",
            "KNOWLEDGE_ARTIFACT_ROOT": str(artifacts),
            "KNOWLEDGE_WORKER_TOKEN": processor_token,
        }
        api_env = {
            **common,
            "KNOWLEDGE_STATE_DIR": str(state),
            "KNOWLEDGE_CLIENTS_FILE": str(state / "clients.json"),
            "KNOWLEDGE_CLIENTS_B64": "",
            "KNOWLEDGE_PORT": str(api_port),
            "KNOWLEDGE_HOST": "127.0.0.1",
            "KNOWLEDGE_WORKER_URL": f"http://127.0.0.1:{processor_port}",
            "KNOWLEDGE_EMBEDDING_URL": "",
            "KNOWLEDGE_ANSWER_URL": "",
            "KNOWLEDGE_PROVIDER_TOKEN": "",
            "KNOWLEDGE_CONTROL_TOKEN": control_token,
            "DOCLING_FAULT": "",
            "DOCLING_QUALITY_FAULT": "",
        }
        reader, editor = clients
        config = {
            "api_url": f"http://127.0.0.1:{api_port}",
            "public_url": f"http://127.0.0.1:{gateway_port}",
            "client_id": reader["id"],
            "issuer": reader["issuer"],
            "kid": "v1",
            "api_token": reader["token"],
            "signing_key": reader["keys"]["v1"],
            "connector_token": str(uuid4()) + str(uuid4()),
            "api_clients": {
                "editor": {
                    "client_id": editor["id"],
                    "issuer": editor["issuer"],
                    "kid": "v1",
                    "api_token": editor["token"],
                    "signing_key": editor["keys"]["v1"],
                }
            },
            "users": {
                "writer": {
                    "actions": ["read", "write"],
                    "api_client": "editor",
                    "scopes": [{"collection_id": "synthetic"}],
                    "mcp_token": str(uuid4()) + str(uuid4()),
                },
                "reader": {
                    "actions": ["read"],
                    "scopes": [{"collection_id": "synthetic"}],
                    "mcp_token": str(uuid4()) + str(uuid4()),
                },
            },
        }
        config_file = temporary / "gateway.json"
        config_file.write_text(json.dumps(config))
        config_file.chmod(0o600)
        commands = [
            [
                str(ROOT / ".venv/bin/python"),
                "-m",
                "uvicorn",
                "docling_desk.knowledge.worker:app_factory",
                "--factory",
                "--host",
                "127.0.0.1",
                "--port",
                str(processor_port),
                "--no-access-log",
            ],
            ["bun", str(ROOT / "knowledge-api/src/server.ts")],
            [
                str(ROOT / ".venv/bin/python"),
                "-m",
                "uvicorn",
                "docling_desk.viewer.gateway:app_factory",
                "--factory",
                "--host",
                "127.0.0.1",
                "--port",
                str(gateway_port),
                "--no-access-log",
            ],
        ]
        for index, command in enumerate(commands):
            log = (QA / f"acceptance-service-{index}.log").open("w")
            logs.append(log)
            env = api_env if index == 1 else {**common, "VIEWER_CONFIG_FILE": str(config_file)}
            processes.append(subprocess.Popen(command, cwd=ROOT, env=env, stdout=log, stderr=log))
        wait_ready(config["api_url"] + "/health/ready", processes[1])
        wait_ready(config["public_url"] + "/health/ready", processes[2])
        # Processor liveness allows errors about model readiness to be diagnosed separately.
        wait_ready(f"http://127.0.0.1:{processor_port}/health/live", processes[0])
        results = asyncio.run(
            scenario(clients, config, config["api_url"], api_env, processes, logs, temporary)
        )
        (QA / "acceptance-results.json").write_text(
            json.dumps(
                {
                    "date": datetime.now(ZoneInfo("Asia/Tokyo")).date().isoformat(),
                    "real_llm": False,
                    "real_larm": False,
                    "results": results,
                },
                indent=2,
            )
            + "\n"
        )
        print("Isolated HTTP/MCP/browser acceptance passed", flush=True)
    finally:
        for process in reversed(processes):
            if process.poll() is None:
                process.terminate()
                try:
                    process.wait(timeout=20)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait()
        for log in logs:
            log.close()
        shutil.rmtree(temporary)


if __name__ == "__main__":
    main()
