"""Measure real Docling extraction and page-streamed previews, offline.

DOCLING_ENV_FILE=/dev/null .venv/bin/python qa/performance/powerpoint_preview.py \
    --pages 530 --output qa/performance/powerpoint-530
Add --native to export through the platform Office renderer instead of using
an equivalent saved synthetic PDF. Reports are generated outside Git tracking.
"""

from __future__ import annotations

import argparse
import json
import platform
import shutil
import sys
import threading
import time
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "packages/docling-azure-ocr/src"))

import psutil
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from reportlab.pdfgen import canvas

from docling_desk.documents.conversion import Job, convert_job
from docling_desk.preview.editable_preview import build_editable_preview
from docling_desk.preview.hashing import digest
from docling_desk.storage import document_folder, original_file, record_original
from docling_desk.translation.source import file_digest, read_source, source_map


class Measurement:
    def __init__(self, label):
        self.label = label
        self.stop = threading.Event()
        self.process = psutil.Process()
        self.baseline = self.process.memory_info().rss
        self.peak = self.baseline
        self.family_peak = self.baseline
        self.started = time.perf_counter()
        self.thread = threading.Thread(target=self.sample, daemon=True)

    def sample(self):
        while not self.stop.is_set():
            self.peak = max(self.peak, self.process.memory_info().rss)
            children = self.process.children(recursive=True)
            family = self.process.memory_info().rss
            for child in children:
                try:
                    family += child.memory_info().rss
                except psutil.Error:
                    pass
            self.family_peak = max(self.family_peak, family)
            self.stop.wait(0.02)

    def __enter__(self):
        self.thread.start()
        return self

    def __exit__(self, *args):
        self.stop.set()
        self.thread.join()
        self.result = {
            "stage": self.label,
            "seconds": round(time.perf_counter() - self.started, 3),
            "baseline_rss_mib": round(self.baseline / 1024**2, 1),
            "peak_rss_mib": round(self.peak / 1024**2, 1),
            "peak_process_family_rss_mib": round(self.family_peak / 1024**2, 1),
        }
        print(json.dumps(self.result), flush=True)


def generate(source: Path, pdf: Path | None, count: int):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(600), Pt(300)
    writer = canvas.Canvas(str(pdf), pagesize=(600, 300)) if pdf else None
    for n in range(1, count + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = f"Slide {n}: office file flow affinity"
        shape = slide.shapes.add_textbox(Pt(30), Pt(20), Pt(540), Pt(45))
        shape.text = title
        shape.text_frame.paragraphs[0].font.size = Pt(24)
        body = "\n".join(
            f"Line {i}: sufficient efficient workflow and financial figures." for i in range(1, 9)
        )
        shape = slide.shapes.add_textbox(Pt(30), Pt(80), Pt(410), Pt(170))
        shape.text = body
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.size = Pt(13)
        image = Image.new("RGB", (96, 96))
        image.putdata(
            [
                ((x * 7 + n) % 256, (y * 5 + n) % 256, (x * y + n) % 256)
                for y in range(96)
                for x in range(96)
            ]
        )
        png = BytesIO()
        image.save(png, format="PNG")
        slide.shapes.add_picture(BytesIO(png.getvalue()), Pt(470), Pt(110), Pt(96), Pt(96))
        for i in range(20):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Pt(20 + i * 28), Pt(270), Pt(20), Pt(6)
            )
            shape.fill.solid()
        if writer:
            writer.setFont("Helvetica", 24)
            writer.drawString(30, 255, title)
            writer.setFont("Helvetica", 13)
            for i, line in enumerate(body.splitlines()):
                writer.drawString(30, 208 - i * 18, line)
            from reportlab.lib.utils import ImageReader

            writer.drawImage(ImageReader(BytesIO(png.getvalue())), 470, 94, 96, 96)
            for i in range(20):
                writer.rect(20 + i * 28, 24, 20, 6, fill=1)
            writer.showPage()
    deck.save(source)
    if writer:
        writer.save()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pages", type=int, default=530)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--native", action="store_true")
    parser.add_argument(
        "--source",
        type=Path,
        help="Existing PPTX; unchanged original is copied into isolated storage",
    )
    parser.add_argument("--saved-pdf", type=Path, help="Saved PDF accompanying --source")
    args = parser.parse_args()
    output = args.output.resolve()
    root = output / "data"
    if root.exists():
        raise ValueError("Use an empty output directory; benchmark never overwrites existing data")
    folder = document_folder(root, "a" * 32)
    folder.mkdir(parents=True)
    source = root / "content/documents" / folder.name / "original.pptx"
    source.parent.mkdir(parents=True)
    pdf = folder / "powerpoint-rendered.pdf"
    if args.source:
        shutil.copyfile(args.source, source)
        if args.saved_pdf:
            shutil.copyfile(args.saved_pdf, pdf)
    else:
        generate(source, None if args.native else pdf, args.pages)
    if pdf.exists():
        (folder / "powerpoint-export.json").write_text(
            json.dumps(
                {
                    "renderer": "Saved PDF benchmark",
                    "source_sha256": digest(source),
                    "pdf_sha256": digest(pdf),
                }
            )
        )
    job = Job(id=folder.name, filename="benchmark.pptx", ocr_profile={"provider": "disabled"})
    record_original(folder, job.filename, job.created)
    stages = []
    with Measurement("Docling conversion + preview + slide export") as measured:
        convert_job(folder, job)
    stages.append(measured.result)
    assert job.state in {"success", "partial"}, job.error
    assert job.pages == args.pages, (job.pages, args.pages)
    assert job.preview, job.warnings
    for n in range(1, args.pages + 1):
        assert (folder / f"editable-preview/page-{n}.html").is_file()
    with Measurement("preview rebuild from saved PDF") as measured:
        build_editable_preview(original_file(folder, ".pptx"), pdf, folder)
    stages.append(measured.result)
    file_digest.cache_clear()
    read_source.cache_clear()
    with Measurement("first translation binding") as measured:
        first = source_map(folder)
    stages.append(measured.result)
    with patch(
        "docling_desk.translation.source.build_source",
        side_effect=AssertionError("warm binding re-parsed pages"),
    ):
        with Measurement("warm translation binding (10 reads)") as measured:
            for _ in range(10):
                assert source_map(folder) == first
    stages.append(measured.result)
    report = {
        "platform": platform.platform(),
        "python": platform.python_version(),
        "pages": job.pages,
        "renderer": "platform Office"
        if args.native
        else "saved synthetic PDF"
        if not args.source
        else "saved PDF",
        "ocr": "disabled (text/graphics workload)",
        "source_bytes": source.stat().st_size,
        "pdf_bytes": pdf.stat().st_size,
        "all_page_files_verified": args.pages,
        "translation_units": len(first["units"]),
        "stages": stages,
        "memory_note": "RSS sampled every 20ms. Process-family includes child renderers; macOS PowerPoint runs as an independent app and is excluded. Not an Azure container measurement.",
    }
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps(report, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
