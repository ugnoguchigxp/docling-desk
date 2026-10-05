"""Synthetic fixture preparation; never reads the application's live data/."""

import shutil
import sys
from pathlib import Path
from uuid import uuid4

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from docling_desk.documents.conversion import Job, convert_job  # noqa: E402

base = ROOT / ".cache/viewer-embedding/fixtures"
base.mkdir(parents=True, exist_ok=True)
for kind in ["slide", "sheet", "page"]:
    folder = base / kind
    shutil.copytree(ROOT / "tests/fixtures/documents" / kind, folder, dirs_exist_ok=True)
for suffix in ["docx", "md", "txt"]:
    folder = base / suffix
    folder.mkdir(exist_ok=True)
    source = folder / f"original.{suffix}"
    if suffix == "docx":
        doc = Document()
        doc.add_heading("Synthetic viewer document", 0)
        doc.add_paragraph("Synthetic inventory evidence 120 items.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Count"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "120"
        doc.save(source)
    else:
        source.write_text(
            "# Synthetic inventory\n\nSynthetic inventory evidence 120 items.\n\n| Item | Count |\n| --- | --- |\n| A | 120 |\n"
            if suffix == "md"
            else "Synthetic inventory evidence 120 items.\nLiteral text <script> stays text.\n"
        )
    job = Job(id=str(uuid4()), filename=f"synthetic.{suffix}")
    convert_job(folder, job)
    if job.state != "success":
        raise RuntimeError(job.error)
print("Prepared six synthetic formats")
