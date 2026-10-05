"""Write shareable Office samples for synthetic browser checks. No user documents."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Emu, Pt


def write_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Emu(864 * 12700)
    deck.slide_height = Emu(486 * 12700)
    blank = deck.slide_layouts[6]
    first = deck.slides.add_slide(blank)
    box = first.shapes.add_textbox(
        Emu(36 * 12700), Emu(40 * 12700), Emu(400 * 12700), Emu(80 * 12700)
    )
    box.text_frame.paragraphs[0].text = "PT.合成"
    box.text_frame.paragraphs[0].font.size = Pt(28)
    second = deck.slides.add_slide(blank)
    other = second.shapes.add_textbox(
        Emu(36 * 12700), Emu(40 * 12700), Emu(400 * 12700), Emu(80 * 12700)
    )
    other.text_frame.paragraphs[0].text = "2枚目"
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)


def write_workbook(path: Path) -> None:
    book = Workbook()
    first = book.active
    first.title = "集計"
    first.append(["区分", "件数"])
    first.append(["A", 120])
    second = book.create_sheet("注記")
    second.append(["注記"])
    second.append(["合成"])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
