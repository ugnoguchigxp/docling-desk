from __future__ import annotations

import io
import json
from pathlib import Path

import httpx
import pytest
from docling_azure_ocr.config import OcrError, OcrProfile
from docling_azure_ocr.runtime import OcrRuntime
from docling_azure_ocr.transport import AzureTransport
from PIL import Image, ImageDraw
from pydantic import SecretStr

from docling_desk.documents.conversion import job_converter
from docling_desk.documents.ocr_operations import FileOcrLedger
from docling_desk.documents.ocr_profiles import enrich_office_images, enrich_pdf_evidence
from docling_desk.documents.rag import export_rag

OP = "https://test.cognitiveservices.azure.com/documentintelligence/documentModels/prebuilt-read/analyzeResults/11111111-1111-1111-1111-111111111111?api-version=2024-11-30"
PROFILE = OcrProfile(
    provider="azure_read",
    enabled=True,
    endpoint="https://test.cognitiveservices.azure.com",
    auth="api_key",
)


def response(size, text="IMAGE NEEDLE 73921"):
    w, h = size
    return {
        "status": "succeeded",
        "analyzeResult": {
            "apiVersion": "2024-11-30",
            "modelId": "prebuilt-read",
            "pages": [
                {
                    "pageNumber": 1,
                    "unit": "pixel",
                    "width": w,
                    "height": h,
                    "words": [
                        {
                            "content": text,
                            "polygon": [10, 10, w - 10, 10, w - 10, h - 10, 10, h - 10],
                        }
                    ],
                }
            ],
        },
    }


class RecordingTransport:
    def __init__(self, error=None):
        self.posts = self.gets = 0
        self.error = error
        self.size = (400, 100)

    def headers(self):
        return {}

    def submit(self, png, remaining):
        self.posts += 1
        self.size = Image.open(io.BytesIO(png)).size
        if self.error:
            raise OcrError(self.error)
        return OP, "request-test"

    def poll(self, operation, remaining):
        assert operation == OP
        self.gets += 1
        return response(self.size), 0

    def close(self):
        pass


def runtime(folder, transport=None, profile=PROFILE, **kwargs):
    return OcrRuntime(
        profile,
        FileOcrLedger(folder / "operations", profile.max_submissions),
        folder / "results",
        "source-a",
        "a" * 64,
        transport=transport,
        **kwargs,
    )


def test_cache_reuses_response_and_preserves_missing_confidence(tmp_path):
    transport = RecordingTransport()
    value = runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    assert value["words"][0]["confidence"] is None
    assert transport.posts == transport.gets == 1
    cached = runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    assert cached == value
    assert transport.posts == transport.gets == 1
    Path(tmp_path / "results" / (value["cache_key"] + ".json")).write_text("{}")
    with pytest.raises(OcrError, match="ocr_cache_invalid"):
        runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    assert transport.posts == 1


@pytest.mark.parametrize("state", ["submitting", "submission_unknown"])
def test_uncertain_submission_stops_without_second_post(tmp_path, state):
    transport = RecordingTransport("submission_unknown")
    with pytest.raises(OcrError, match="submission_unknown"):
        runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    path = next((tmp_path / "operations").glob("*.json"))
    record = json.loads(path.read_text())
    record["state"] = state
    path.write_text(json.dumps(record))
    transport.error = None
    with pytest.raises(OcrError) as failed:
        runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    assert failed.value.code == "submission_unknown"
    assert transport.posts == 1 and transport.gets == 0
    runtime(tmp_path, transport, allow_resubmit=True).read(
        Image.new("RGB", (400, 100)), "image:1", {}
    )
    assert transport.posts == 2


def test_restart_resumes_known_operation_by_get(tmp_path):
    transport = RecordingTransport()
    transport.poll = lambda *args: (_ for _ in ()).throw(OcrError("ocr_page_timeout"))
    with pytest.raises(OcrError, match="ocr_page_timeout"):
        runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    transport.poll = lambda operation, remaining: (response((400, 100)), 0)
    runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    assert transport.posts == 1


def test_transport_never_retries_post_or_follows_foreign_location():
    calls = []

    def handle(request):
        calls.append(request)
        return httpx.Response(202, headers={"operation-location": OP.replace("test.", "foreign.")})

    transport = AzureTransport(
        PROFILE, SecretStr("secret"), httpx.Client(transport=httpx.MockTransport(handle))
    )
    with pytest.raises(OcrError, match="submission_unknown"):
        transport.submit(b"png", 10)
    assert len(calls) == 1
    with pytest.raises(OcrError):
        transport.poll(OP.replace("test.", "foreign."), 10)
    assert len(calls) == 1


def test_disabled_and_budget_block_external_requests(tmp_path):
    transport = RecordingTransport()
    with pytest.raises(OcrError, match="ocr_disabled"):
        runtime(tmp_path, transport, profile=OcrProfile(provider="disabled")).read(
            Image.new("RGB", (400, 100)), "x", {}
        )
    assert transport.posts == 0
    instance = runtime(
        tmp_path, transport, profile=PROFILE.model_copy(update={"max_submissions": 1})
    )
    instance.read(Image.new("RGB", (400, 100)), "x", {})
    with pytest.raises(OcrError, match="ocr_submission_limit"):
        instance.read(Image.new("RGB", (400, 100)), "y", {})
    assert transport.posts == 1


def office_file(path, kind, image_path, table=False):
    if kind == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), width=Inches(4))
        doc.save(path)
    elif kind == "docx":
        from docx import Document
        from docx.shared import Inches

        doc = Document()
        target = doc.add_table(rows=1, cols=1).cell(0, 0) if table else doc
        target.add_paragraph().add_run().add_picture(str(image_path), width=Inches(4))
        doc.save(path)
    else:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage

        doc = Workbook()
        doc.active.title = "Evidence Sheet"
        doc.active.add_image(SheetImage(str(image_path)), "B2")
        doc.save(path)


@pytest.mark.parametrize(
    "kind,table", [("pptx", False), ("docx", False), ("docx", True), ("xlsx", False)]
)
def test_real_office_images_enter_rag_and_roundtrip(tmp_path, kind, table):
    from docling_core.types.doc import DoclingDocument

    image = tmp_path / "image.png"
    bitmap = Image.new("RGB", (400, 100), "white")
    ImageDraw.Draw(bitmap).text((15, 15), "IMAGE NEEDLE 73921", fill="black")
    bitmap.save(image)
    source = tmp_path / ("original." + kind)
    office_file(source, kind, image, table)
    transport = RecordingTransport()
    instance = runtime(tmp_path, transport)
    doc = job_converter(PROFILE, instance).convert(source).document
    assert doc.pictures
    enrich_office_images(doc, instance)
    # Metadata must survive saved docling JSON, not only exist on the live Python objects.
    doc = DoclingDocument.model_validate_json(doc.model_dump_json())
    export_rag(doc, source.name, "job-test", source, tmp_path)
    chunks = [json.loads(line) for line in (tmp_path / "rag-index.jsonl").read_text().splitlines()]
    for name in ("rag.jsonl", "rag-index.jsonl", "rag-docling.jsonl"):
        records = [json.loads(line) for line in (tmp_path / name).read_text().splitlines()]
        assert all("response_sha256" not in record["text"] for record in records)
    matching = [c for c in chunks if "IMAGE NEEDLE 73921" in c["text"]]
    assert matching, chunks
    evidence = [e for c in matching for e in c["ocr_evidence"]]
    assert evidence and evidence[0]["picture_ref"].startswith("#/pictures/")
    assert evidence[0]["words"][0]["confidence"] is None
    assert evidence[0]["source_sha256"] == "a" * 64
    if kind in ("pptx", "xlsx"):
        assert matching[0]["pages"] == [1]
    else:
        assert matching[0]["pages"] == []
    enrich_office_images(job_converter(PROFILE, instance).convert(source).document, instance)
    assert transport.posts == 1


def test_ocr_evidence_is_separate_from_exported_text_and_html(tmp_path, monkeypatch):
    import docling_desk.documents.conversion as service

    image = tmp_path / "image.png"
    bitmap = Image.new("RGB", (400, 100), "white")
    ImageDraw.Draw(bitmap).text((15, 15), "IMAGE NEEDLE 73921", fill="black")
    bitmap.save(image)
    office_file(tmp_path / "original.docx", "docx", image)
    monkeypatch.setattr(service, "office_preview", lambda *_: ("extracted.html", "Test preview"))
    job = service.Job(
        id="synthetic-export", filename="synthetic-export.docx", ocr_profile=PROFILE.model_dump()
    )
    service.convert_job(tmp_path, job, ocr_runtime=runtime(tmp_path, RecordingTransport()))
    assert job.state == "success", job.error
    for name in ("text.txt", "document.md", "extracted.html"):
        exported = (tmp_path / name).read_text()
        assert "IMAGE NEEDLE 73921" in exported
        assert "response_sha256" not in exported
        assert "cognitiveservices.azure.com" not in exported
    saved = (tmp_path / "document.json").read_text()
    assert "response_sha256" in saved
    records = [json.loads(line) for line in (tmp_path / "rag-index.jsonl").read_text().splitlines()]
    assert any(record["ocr_evidence"] for record in records)
    assert all("response_sha256" not in record["text"] for record in records)


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_real_pdf_adapter_preserves_backend_coordinate_frame(tmp_path, rotation):
    import pymupdf

    from docling_desk.documents.rag import build_context_chunks

    image = io.BytesIO()
    bitmap = Image.new("RGB", (400, 100), "white")
    ImageDraw.Draw(bitmap).text((15, 15), "IMAGE NEEDLE 73921", fill="black")
    bitmap.save(image, "PNG")
    pdf = pymupdf.open()
    page = pdf.new_page(width=500, height=300)
    page.insert_image(pymupdf.Rect(50, 50, 450, 150), stream=image.getvalue())
    page.set_cropbox(pymupdf.Rect(25, 25, 475, 275))
    page.set_rotation(rotation)
    source = tmp_path / "original.pdf"
    pdf.save(source)
    pdf.close()
    transport = RecordingTransport()
    profile = PROFILE.model_copy(update={"profile": "read-full-page-v1"})
    instance = runtime(tmp_path, transport, profile=profile)
    doc = job_converter(profile, instance).convert(source).document
    assert transport.posts == 1
    enrich_pdf_evidence(doc, instance)
    value = next(iter(instance.evidence.values()))
    assert value["page_number"] == 1
    assert len(value["image_to_page"]) == 6
    assert value["image_to_page"][0] * value["image_size"][0] == pytest.approx(
        doc.pages[1].size.width
    )
    assert value["image_to_page"][4] * value["image_size"][1] == pytest.approx(
        doc.pages[1].size.height
    )
    saved = json.loads(next((tmp_path / "operations").glob("*.json")).read_text())
    assert saved["metadata"]["image_to_page"] == value["image_to_page"]
    assert value["coordinate_frame"] == "docling_backend_page_top_left"
    chunks = build_context_chunks(doc, source.name, "a" * 64, "job")[1]
    assert any("IMAGE NEEDLE 73921" in c.text for c in chunks), doc.export_to_markdown()
    assert any(c.ocr_evidence for c in chunks)
    assert all(c.pages == [1] for c in chunks)


def test_real_local_office_worker_ocr_is_returned_to_rag(tmp_path, monkeypatch):
    import hashlib
    import subprocess
    import sys
    from uuid import uuid4

    from fastapi.testclient import TestClient
    from PIL import ImageFont

    import docling_desk.documents.conversion as service
    from docling_desk.knowledge.worker import create_worker

    if sys.platform != "darwin":
        import shutil

        if not shutil.which("tesseract"):
            pytest.skip("Local OCR engine not installed")
    source_id, revision = str(uuid4()), str(uuid4())
    folder = tmp_path / "sources" / source_id / revision
    folder.mkdir(parents=True)
    image_path = tmp_path / "readable.png"
    bitmap = Image.new("RGB", (800, 200), "white")
    font_path = (
        "/System/Library/Fonts/Supplemental/Arial.ttf"
        if sys.platform == "darwin"
        else subprocess.check_output(
            ["fc-match", "-f", "%{file}", "Liberation Sans"], text=True
        ).strip()
    )
    ImageDraw.Draw(bitmap).text(
        (40, 60), "IMAGE NEEDLE 73921", font=ImageFont.truetype(font_path, 48), fill="black"
    )
    bitmap.save(image_path)
    source = folder / "original.docx"
    office_file(source, "docx", image_path, table=True)
    monkeypatch.setattr(
        service, "office_preview", lambda *_: ("extracted.html", "Synthetic preview")
    )
    token = "private-test-token-" * 3
    client = TestClient(create_worker(tmp_path, token))
    result = client.post(
        "/internal/v1/extract",
        headers={"Authorization": "Bearer " + token},
        json={
            "source_id": source_id,
            "source_revision": revision,
            "suffix": ".docx",
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "job_id": str(uuid4()),
            "run": 1,
            "profile": "local-v1",
            "filename": "images.docx",
        },
    )
    assert result.status_code == 200, result.text
    records = result.json()["chunks"]
    assert any("NEEDLE" in c["text"] for c in records)
    assert any(c["ocr_evidence"] for c in records)
    assert all(c["pages"] == [] for c in records)


def test_operation_state_unavailable_never_posts_and_delete_fences_cache(tmp_path):
    transport = RecordingTransport()
    instance = runtime(tmp_path, transport)
    instance.ledger.live = lambda: False
    with pytest.raises(OcrError, match="ocr_cancelled"):
        instance.read(Image.new("RGB", (400, 100)), "image:1", {})
    assert transport.posts == 0
    instance.ledger.live = lambda: True
    instance.read(Image.new("RGB", (400, 100)), "image:1", {})
    instance.ledger.live = lambda: False
    with pytest.raises(OcrError, match="ocr_cancelled"):
        instance.read(Image.new("RGB", (400, 100)), "image:1", {})
    assert transport.posts == 1


def test_minimal_connection_keeps_local_default_and_never_serializes_credentials():
    from docling_azure_ocr.config import load_profile
    from pydantic import ValidationError

    env = {
        "DOCLING_AZURE_OCR_ENDPOINT": PROFILE.endpoint,
        "AZURE_DOCUMENT_INTELLIGENCE_API_KEY": "secret",
    }
    with pytest.raises(ValidationError):
        load_profile("azure_read", env={})
    assert load_profile(env=env).provider == "local"
    profile = load_profile("azure_read", env=env)
    assert profile.enabled is True
    assert profile.auth == "api_key"
    assert "secret" not in profile.model_dump_json()
    with pytest.raises(OcrError) as failed:
        profile.check_credentials(SecretStr(""))
    assert failed.value.code == "ocr_auth_missing"
    assert load_profile(env={}).provider == "local"


def test_managed_identity_needs_only_endpoint_and_optional_client_id():
    from docling_azure_ocr.config import load_profile

    env = {"DOCLING_AZURE_OCR_ENDPOINT": PROFILE.endpoint}
    profile = load_profile("azure_read", env=env)
    assert profile.auth == "managed_identity"
    assert profile.client_id is None
    assert (
        load_profile(
            "azure_read", env={**env, "DOCLING_AZURE_OCR_CLIENT_ID": "assigned-id"}
        ).client_id
        == "assigned-id"
    )


def test_removed_environment_options_cannot_change_ocr_defaults():
    from docling_azure_ocr.config import load_profile

    env = {
        "DOCLING_AZURE_OCR_ENDPOINT": PROFILE.endpoint,
        "AZURE_DOCUMENT_INTELLIGENCE_API_KEY": "secret",
        "DOCLING_OCR_PROVIDER": "azure_read",
        "DOCLING_AZURE_OCR_ENABLED": "false",
        "DOCLING_AZURE_OCR_AUTH": "managed_identity",
        "DOCLING_AZURE_OCR_API_VERSION": "unsupported-version",
        "DOCLING_AZURE_OCR_MODEL": "unsupported-model",
        "DOCLING_AZURE_OCR_PROFILE": "read-full-page-v1",
        "DOCLING_AZURE_OCR_TIER": "S0",
        "DOCLING_AZURE_OCR_DOCUMENT_TIMEOUT": "garbage",
        "DOCLING_AZURE_OCR_MAX_SUBMISSIONS": "garbage",
    }
    profile = load_profile(env=env)
    assert profile.provider == "local"
    assert profile.enabled is True
    assert profile.auth == "api_key"
    assert (profile.api_version, profile.model, profile.profile, profile.tier) == (
        "2024-11-30",
        "prebuilt-read",
        "read-v1",
        "F0",
    )
    assert (profile.document_timeout, profile.max_submissions) == (1800, 100)


def test_native_text_pdf_does_not_submit_to_azure(tmp_path):
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((60, 60), "NATIVE TEXT ONLY - no image OCR needed")
    source = tmp_path / "native.pdf"
    pdf.save(source)
    pdf.close()
    transport = RecordingTransport()
    instance = runtime(tmp_path, transport)
    doc = job_converter(PROFILE, instance).convert(source).document
    assert "NATIVE TEXT ONLY" in doc.export_to_text()
    assert transport.posts == 0


def test_completed_missing_cache_needs_explicit_resend(tmp_path):
    transport = RecordingTransport()
    value = runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    (tmp_path / "results" / (value["cache_key"] + ".json")).unlink()
    with pytest.raises(OcrError) as failed:
        runtime(tmp_path, transport).read(Image.new("RGB", (400, 100)), "image:1", {})
    assert failed.value.code == "ocr_cache_missing"
    assert transport.posts == 1
    runtime(tmp_path, transport, allow_resubmit=True).read(
        Image.new("RGB", (400, 100)), "image:1", {}
    )
    assert transport.posts == 2


def test_embedded_image_text_preserves_japanese_lines(tmp_path):
    transport = RecordingTransport()
    result = response((400, 100), text="日本語の画像")
    result["analyzeResult"]["pages"][0]["lines"] = [
        {"content": "日本語の画像"},
        {"content": "確認番号73921"},
    ]
    transport.poll = lambda *_: (result, 0)
    value = runtime(tmp_path, transport).read(
        Image.new("RGB", (400, 100)), "image:1", {"coordinate_frame": "embedded_image_pixels"}
    )
    assert value["text"] == "日本語の画像\n確認番号73921"
