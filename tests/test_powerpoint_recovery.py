"""Regression cases for ligatures, missing previews, portable bindings and repairs."""

from __future__ import annotations

import json
import math
import os
import re
import time
from pathlib import Path
from unittest.mock import patch

import pytest
import reportlab
from fastapi.testclient import TestClient
from lxml import html
from pptx import Presentation
from pptx.util import Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as ReportlabFont
from reportlab.pdfgen import canvas
from test_blob_mirror import MemoryStore

from docling_desk import config
from docling_desk.blob_mirror import Mirror
from docling_desk.preview.editable_preview import build_editable_preview
from docling_desk.preview.hashing import digest
from docling_desk.preview.preview_migration import rebuild_preview
from docling_desk.storage import document_folder, job_file, migrate_layout, original_file
from docling_desk.translation.source import html_segments, source_map
from docling_desk.translation.view import apply_text

TEXT = "office fi fl ff ffi ffl finish"


def expand_unicode(path: Path, mapping: dict[bytes, str]) -> None:
    """Give glyphs a multi-character ToUnicode mapping without moving any byte offset.

    The PDF is written uncompressed. Unused trailing CMap entries are dropped and the
    difference padded, so every xref offset stays valid.
    """
    data = path.read_bytes()
    gained = 0
    for code, text in mapping.items():
        target = text.encode("utf-16-be").hex().upper().encode()
        data, count = re.subn(
            rb"<" + code + rb"> <[0-9A-F]{4}>\n", b"<" + code + b"> <" + target + b">\n", data
        )
        assert count == 1
        gained += len(target) - 4
    entries = math.ceil(gained / 12)
    for _ in range(entries):
        data, count = re.subn(rb"<[0-9A-F]{2}> <[0-9A-F]{4}>\n(?=endbfchar)", b"", data)
        assert count == 1
    data, count = re.subn(
        rb"(\d+) beginbfchar", lambda m: b"%d beginbfchar" % (int(m[1]) - entries), data
    )
    assert count == 1
    path.write_bytes(data.replace(b"endbfchar", b" " * (entries * 12 - gained) + b"endbfchar", 1))


def ligature_sources(folder: Path):
    source, pdf = folder / "original.pptx", folder / "powerpoint-rendered.pdf"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(600), Pt(300)
    for _ in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Pt(50), Pt(30), Pt(500), Pt(50)).text = TEXT
    deck.save(source)
    font = Path(reportlab.__file__).parent / "fonts/Vera.ttf"
    pdfmetrics.registerFont(ReportlabFont("RecoveryVera", str(font)))
    writer = canvas.Canvas(str(pdf), pagesize=(600, 300), pageCompression=0)
    for _ in range(2):
        writer.setFont("RecoveryVera", 20)
        writer.drawString(50, 250, "office \ufb01 \ufb02 X Y Z finish")
        writer.showPage()
    writer.save()
    size = pdf.stat().st_size
    # Genuine fi/fl glyphs whose ToUnicode expands to several characters; ff/ffi/ffl
    # exercise longer expansions on glyphs that are single letters in the font.
    expand_unicode(pdf, {b"01": "fi", b"02": "fl", b"58": "ff", b"59": "ffi", b"5A": "ffl"})
    assert pdf.stat().st_size == size
    (folder / "powerpoint-export.json").write_text(
        json.dumps({"source_sha256": digest(source), "pdf_sha256": digest(pdf)})
    )
    return source, pdf


def test_ligatures_keep_full_selectable_unicode_and_translation_binding(tmp_path):
    source, pdf = ligature_sources(tmp_path)
    build_editable_preview(source, pdf, tmp_path)
    raw = (tmp_path / "editable-preview/page-2.html").read_text()
    tree = html.fromstring(raw)
    # The expanded characters are all selectable, and the line still belongs to its frame.
    assert tree.xpath('//div[contains(@class,"pdf-text")]/text()') == [TEXT]
    assert tree.xpath('//div[contains(@class,"pdf-text")]/@data-source-id') == ["shape-2"]
    segments = html_segments(raw, [])
    assert [s["source_text"] for s in segments] == [TEXT]
    translated = apply_text(raw, {"segments": segments}, {segments[0]["id"]: "合字の訳文"})
    block = html.fromstring(translated).xpath('//div[contains(@class,"pdf-text")]')[0]
    assert block.text == "合字の訳文" and block.get("data-translated") == "true"


def test_unexpanded_ligature_glyphs_are_read_as_plain_letters(tmp_path):
    pdf = tmp_path / "original.pdf"
    font = Path(reportlab.__file__).parent / "fonts/Vera.ttf"
    pdfmetrics.registerFont(ReportlabFont("RecoveryVeraPlain", str(font)))
    writer = canvas.Canvas(str(pdf), pagesize=(600, 300))
    writer.setFont("RecoveryVeraPlain", 20)
    writer.drawString(50, 250, "of\ufb01ce \ufb02ag")
    writer.save()
    build_editable_preview(pdf, pdf, tmp_path)
    tree = html.fromstring((tmp_path / "editable-preview/page-1.html").read_text())
    assert tree.xpath('//div[contains(@class,"pdf-text")]/text()') == ["office flag"]


def test_preview_failure_keeps_previous_complete_generation(tmp_path):
    source, pdf = ligature_sources(tmp_path)
    build_editable_preview(source, pdf, tmp_path)
    previous = (tmp_path / "editable-preview/manifest.json").read_bytes()
    with patch(
        "docling_desk.preview.editable_preview.page_blocks", side_effect=ValueError("bad page")
    ):
        with pytest.raises(ValueError):
            build_editable_preview(source, pdf, tmp_path)
    assert (tmp_path / "editable-preview/manifest.json").read_bytes() == previous
    assert not list(tmp_path.glob(".editable-preview-*"))


def test_null_preview_api_uses_extracted_text_and_returns_reasons(
    translation_document, monkeypatch
):
    from conftest import FixedProvider, wait_translation

    import docling_desk.app as web
    from docling_desk.translation import service

    folder = translation_document()
    job = json.loads(job_file(folder).read_text())
    job["preview"] = None
    job_file(folder).write_text(json.dumps(job))
    slides = json.loads((folder / "slides.json").read_text())
    for slide in slides["slides"]:
        slide["preview"] = None
    (folder / "slides.json").write_text(json.dumps(slides))
    monkeypatch.setattr(config, "DATA", folder.parent)
    monkeypatch.setattr(service, "provider_for", lambda profile: FixedProvider())
    with TestClient(web.app) as client:
        folder = document_folder(config.DATA, folder.name)
        url = f"/api/jobs/{folder.name}/translations"
        response = client.get(url)
        assert response.status_code == 200
        units = response.json()["units"]
        assert all(u["mode"] == "panel" and u["segments_count"] for u in units)
        assert all("原本プレビュー" in u["preview_unavailable_reason"] for u in units)
        assert (
            client.post(url, json={"target_language": "en", "unit_ids": ["slide-1"]}).status_code
            == 202
        )
        assert wait_translation(folder, "en", "slide-1")["state"] == "completed"
        result = client.get(url + "/en/slide-1").json()
        assert result["mode"] == "panel" and result["texts"]


def test_empty_missing_preview_reports_unavailable_reason(translation_document):
    folder = translation_document()
    slides = json.loads((folder / "slides.json").read_text())
    slides["slides"].append({"number": 99, "preview": None})
    (folder / "slides.json").write_text(json.dumps(slides))
    unit = source_map(folder)["units"][-1]
    assert unit["segments"] == [] and unit["unavailable_reason"]


def test_warm_status_skips_document_and_template_parsing(translation_document):
    folder = translation_document()
    first = source_map(folder)
    with (
        patch(
            "docling_desk.translation.source.build_source", side_effect=AssertionError("must reuse")
        ),
        patch(
            "docling_desk.translation.source.digest", side_effect=AssertionError("unchanged bytes")
        ),
    ):
        assert source_map(folder) == first


@pytest.mark.parametrize(
    "change", ["original", "template", "reference", "nested_reference", "document", "generation"]
)
def test_binding_invalidates_changed_inputs(translation_document, change, monkeypatch):
    import docling_desk.translation.source as sources

    folder = translation_document()
    first = source_map(folder)
    if change == "original":
        original = original_file(folder, ".pptx")
        deck = Presentation(original)
        deck.slides[0].shapes[0].text = "changed source"
        deck.save(original)
    elif change == "template":
        slide = json.loads((folder / "slides.json").read_text())["slides"][0]
        path = folder / slide["preview"]
        path.write_text(path.read_text() + "<!-- changed -->")
    elif change == "reference":
        (folder / "quicklook/new-font.woff").write_bytes(b"new reference")
    elif change == "nested_reference":
        path = folder / "references/legend.png"
        path.parent.mkdir()
        path.write_bytes(b"new nested reference")
    elif change == "document":
        path = folder / "document.json"
        path.write_text(path.read_text() + "\n")
    else:
        monkeypatch.setattr(sources, "BINDING_VERSION", sources.BINDING_VERSION + 1)
    second = source_map(folder)
    assert first["dependencies"] != second["dependencies"]
    assert first["units"][0]["source_hash"] != second["units"][0]["source_hash"]


def test_repair_and_blob_restore_reuse_saved_extraction_and_preserve_metadata(
    translation_document, tmp_path, monkeypatch
):
    folder = translation_document()
    root = folder.parent
    # Replace only test originals/rendered PDF, keep the saved Docling data.
    ligature_sources(folder)
    job = json.loads(job_file(folder).read_text())
    job["preview"] = None
    job["folder_id"] = "original-category"
    job_file(folder).write_text(json.dumps(job))
    (folder / "translations/en").mkdir(parents=True)
    saved_translation = folder / "translations/en/slide-1.json"
    saved_translation.write_text(
        '{"unit_id":"slide-1","state":"completed","result":{"translations":{"t00001":"saved text"}}}'
    )
    (root / "library.json").write_text('{"folders":[{"id":"original-category","name":"分類"}]}')
    migrate_layout(root)
    folder = document_folder(root, folder.name)
    protected = {
        p: p.read_bytes()
        for p in [
            folder / "document.json",
            folder / "rag.jsonl",
            folder / "elements.json",
            folder / "translations/en/slide-1.json",
            root / "content/manifests/library.json",
        ]
    }
    with (
        patch(
            "docling_desk.preview.powerpoint_export.export_powerpoint",
            side_effect=AssertionError("no renderer"),
        ),
        patch(
            "docling_desk.documents.conversion.convert_job",
            side_effect=AssertionError("no extraction"),
        ),
    ):
        assert rebuild_preview(folder, saved_pdf_only=True)["state"] == "updated"
    assert all(path.read_bytes() == content for path, content in protected.items())
    assert json.loads(job_file(folder).read_text())["folder_id"] == "original-category"
    assert all(s["preview"] for s in json.loads((folder / "slides.json").read_text())["slides"])
    original_binding = source_map(folder)
    store = MemoryStore()
    for path in root.rglob("*"):
        if path.is_file():
            old = time.time() - 60
            os.utime(path, (old, old))
    report = Mirror(root, store, derived=True).push()
    assert not report.errors
    for name in (
        "document.json",
        "slides.json",
        "powerpoint-rendered.pdf",
        "translation-source.json",
        "editable-preview/page-2.html",
    ):
        assert f"derived/documents/{folder.name}/{name}" in store.blobs
    restored_root = tmp_path / "empty-restored"
    assert Mirror(restored_root, store, derived=True).pull().downloaded > 0
    restored = document_folder(restored_root, folder.name)
    from docling_desk.translation import source as sources

    sources.file_digest.cache_clear()
    sources.read_source.cache_clear()
    with patch.object(
        sources, "build_source", side_effect=AssertionError("Blob binding must be reusable")
    ):
        assert source_map(restored) == original_binding
    assert (restored / "translations/en/slide-1.json").read_bytes() == protected[
        folder / "translations/en/slide-1.json"
    ]
