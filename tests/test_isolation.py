from __future__ import annotations

import time
import zipfile
from pathlib import Path

import httpx
import pytest
from pptx import Presentation

from docling_desk import config as desk_config
from docling_desk.api import diagnostics
from docling_desk.documents import office_isolation
from docling_desk.documents.conversion import validate_input
from docling_desk.knowledge import tokens
from docling_desk.preview import office_linux


def deck(path: Path, text: str = "隔離された抽出") -> Path:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = text
    presentation.save(str(path))
    return path


def test_a_stuck_office_conversion_is_killed_at_the_deadline():
    started = time.monotonic()
    assert office_isolation.run_isolated(time.sleep, (60,), timeout=2) is None
    assert time.monotonic() - started < 30


def test_office_documents_are_extracted_in_a_child_process(tmp_path):
    result = office_isolation.convert_office(deck(tmp_path / "a.pptx"), 2000, 50 * 1024 * 1024)
    assert result.status.value in {"success", "partial_success"}
    assert "隔離された抽出" in result.document.export_to_markdown()


def test_the_office_deadline_is_reported_as_a_failure(tmp_path, monkeypatch):
    monkeypatch.setattr(office_isolation, "run_isolated", lambda *a, **k: None)
    monkeypatch.setenv("DOCLING_OFFICE_TIMEOUT_SECONDS", "7")
    with pytest.raises(ValueError, match="7秒以内"):
        office_isolation.convert_office(deck(tmp_path / "a.pptx"), 10, 10**6)
    monkeypatch.setenv("DOCLING_OFFICE_TIMEOUT_SECONDS", "0")
    with pytest.raises(ValueError):
        office_isolation.office_timeout()


def test_macro_enabled_packages_are_refused_whatever_their_name(tmp_path):
    clean = deck(tmp_path / "clean.pptx")
    validate_input(clean)
    hidden = tmp_path / "hidden.pptx"
    with zipfile.ZipFile(clean) as source, zipfile.ZipFile(hidden, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item))
        target.writestr("ppt/vbaProject.bin", b"\0" * 32)
    with pytest.raises(ValueError, match="マクロ"):
        validate_input(hidden)
    typed = tmp_path / "typed.pptx"
    with zipfile.ZipFile(clean) as source, zipfile.ZipFile(typed, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.presentation.main+xml",
                    b"ms-powerpoint.presentation.macroEnabled.main+xml",
                )
            target.writestr(item, data)
    with pytest.raises(ValueError, match="マクロ"):
        validate_input(typed)


def test_libreoffice_runs_sandboxed_with_no_credentials_and_no_macros(tmp_path, monkeypatch):
    monkeypatch.setenv("LARM_API_TOKEN", "larm-shared-credential-0123456789")
    monkeypatch.setattr(
        office_linux.shutil,
        "which",
        lambda name: {"libreoffice": "/usr/bin/libreoffice", "bwrap": "/usr/bin/bwrap"}.get(name),
    )
    seen = {}

    def run(command, timeout, env=None):
        seen.update(command=command, env=env)
        work = Path(command[command.index("--chdir") + 1])
        profile = work / "profile" / "user" / "registrymodifications.xcu"
        seen["profile"] = profile.read_text()
        (work / "out" / "input.html").write_text("<html></html>")
        return office_linux.subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(office_linux, "run_in_group", run)
    source = tmp_path / "secret-name.docx"
    source.write_bytes(b"x")
    output = office_linux.convert_office(source, tmp_path / "out", "html")
    assert output == tmp_path / "out" / "secret-name.html" and output.is_file()
    command = seen["command"]
    assert command[0] == "/usr/bin/bwrap" and "--unshare-all" in command
    assert "--share-net" not in command
    # The data folder itself is never bound into the sandbox.
    assert str(tmp_path) not in " ".join(command[: command.index("/usr/bin/libreoffice")])
    assert "LARM_API_TOKEN" not in seen["env"] and set(seen["env"]) >= {"PATH", "HOME"}
    assert (
        '"DisableMacrosExecution"' in seen["profile"] and "<value>true</value>" in seen["profile"]
    )


def test_libreoffice_is_refused_without_bubblewrap_outside_a_container(tmp_path, monkeypatch):
    monkeypatch.setattr(office_linux.shutil, "which", lambda name: None)
    monkeypatch.delenv("DOCLING_CONTAINER", raising=False)
    with pytest.raises(RuntimeError, match="bubblewrap"):
        office_linux.sandbox(["soffice"], tmp_path)
    monkeypatch.setenv("DOCLING_CONTAINER", "1")
    assert office_linux.sandbox(["soffice"], tmp_path) == ["soffice"]


def test_diagnostics_require_bubblewrap_for_libreoffice(monkeypatch):
    monkeypatch.delenv("DOCLING_CONTAINER", raising=False)
    monkeypatch.setattr(diagnostics, "_which", lambda name: None if name == "bwrap" else "/x")
    assert diagnostics._libreoffice("libreoffice")["code"] == "bubblewrap_missing"
    monkeypatch.setenv("DOCLING_CONTAINER", "1")
    assert diagnostics._libreoffice("libreoffice")["status"] == "ok"


def test_token_encodings_are_read_only_from_the_pinned_copy(tmp_path, monkeypatch):
    monkeypatch.setattr(desk_config, "MODELS", tmp_path)
    tokens._verified.cache_clear()

    def offline(*args, **kwargs):
        raise AssertionError("an encoding was fetched while serving")

    monkeypatch.setattr("tiktoken.load.read_file", offline)
    with pytest.raises(RuntimeError, match="download_models"):
        tokens.encoding("cl100k_base")
    assert tokens.problems() == ["missing:tiktoken/cl100k_base", "missing:tiktoken/o200k_base"]
    path = tokens.cache_path("cl100k_base")
    path.parent.mkdir(parents=True)
    path.write_bytes(b"tampered")
    with pytest.raises(RuntimeError, match="一致しません"):
        tokens.encoding("cl100k_base")
    tokens._verified.cache_clear()


def test_installing_encodings_refuses_a_wrong_download(tmp_path, monkeypatch):
    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=b"not it"))
    real = httpx.Client
    monkeypatch.setattr(tokens.httpx, "Client", lambda **kw: real(transport=transport, **kw))
    with pytest.raises(RuntimeError, match="pinned hash"):
        tokens.install(tmp_path)
    assert not tokens.cache_path("cl100k_base", tmp_path).exists()
