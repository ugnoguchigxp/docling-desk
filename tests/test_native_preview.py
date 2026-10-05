import hashlib
import json

import pytest
from lxml import html
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from reportlab.pdfgen import canvas

from docling_desk.preview.native_preview import build_native_preview, native_preview
from docling_desk.preview.slides import split_quicklook


def sources(folder, pages=2, pdf_pages=2, size=(960, 540)):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(960), Pt(540)
    for _ in range(pages):
        deck.slides.add_slide(deck.slide_layouts[6])
    source = folder / "original.pptx"
    deck.save(source)
    pdf = folder / "powerpoint-rendered.pdf"
    document = canvas.Canvas(str(pdf), pagesize=size)
    for number in range(pdf_pages):
        document.setFillColorRGB(number, 0, 1)
        document.rect(10, 10, 100, 100, fill=1)
        document.showPage()
    document.save()
    return source, pdf


def test_reference_cache_preserves_sources_and_page_images(tmp_path):
    source, pdf = sources(tmp_path)
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, pdf)]
    path = build_native_preview(source, pdf, tmp_path)
    assert path == native_preview(source, tmp_path)
    frames = split_quicklook(tmp_path / path, tmp_path, 2)
    assert len(frames) == 2
    for number, frame in frames.items():
        tree = html.fromstring((tmp_path / frame).read_text())
        assert tree.xpath("//img/@src") == [f"slide-{number}.png"]
        with Image.open(tmp_path / "native-preview" / f"slide-{number}.png") as image:
            assert image.size == (2880, 1620)
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, pdf)]


@pytest.mark.parametrize("pdf_pages,size", [(1, (960, 540)), (2, (540, 960))])
def test_wrong_pdf_rejected_before_creating_preview(tmp_path, pdf_pages, size):
    source, pdf = sources(tmp_path, pdf_pages=pdf_pages, size=size)
    with pytest.raises(ValueError):
        build_native_preview(source, pdf, tmp_path)
    assert not (tmp_path / "native-preview").exists()


@pytest.mark.parametrize("change", ["source", "pdf", "image", "manifest"])
def test_stale_or_incomplete_cache_is_not_used(tmp_path, change):
    source, pdf = sources(tmp_path)
    build_native_preview(source, pdf, tmp_path)
    if change in {"source", "pdf"}:
        target = source if change == "source" else pdf
        target.write_bytes(target.read_bytes() + b"changed")
    elif change == "image":
        (tmp_path / "native-preview" / "slide-2.png").unlink()
    else:
        (tmp_path / "native-preview" / "manifest.json").write_text("invalid JSON")
    assert native_preview(source, tmp_path) is None


def test_manifest_cannot_use_pdf_outside_job(tmp_path):
    source, pdf = sources(tmp_path)
    build_native_preview(source, pdf, tmp_path)
    path = tmp_path / "native-preview" / "manifest.json"
    manifest = json.loads(path.read_text())
    manifest["pdf"] = "../powerpoint-rendered.pdf"
    path.write_text(json.dumps(manifest))
    assert native_preview(source, tmp_path) is None
