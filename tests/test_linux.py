import json
import subprocess
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from lxml import html
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Pt
from reportlab.pdfgen import canvas

import docling_desk.app as web
import docling_desk.documents.conversion as service
import docling_desk.preview.office_linux as office_linux
import docling_desk.preview.powerpoint_export as powerpoint_export
import docling_desk.preview.thumbnails as thumbnails
from docling_desk import config as desk_config
from docling_desk.preview.sheets import preview_sheets, sheet_html
from docling_desk.translation.source import worksheet_segments


def test_linux_ocr_uses_installed_japanese_and_english(monkeypatch):
    monkeypatch.setattr(service.sys, "platform", "linux")
    service.converter.cache_clear()
    options = service.converter().format_to_options[service.InputFormat.PDF].pipeline_options
    assert isinstance(options.ocr_options, service.TesseractCliOcrOptions)
    assert options.ocr_options.lang == ["jpn", "eng"]
    assert options.enable_remote_services is False
    service.converter.cache_clear()


def test_linux_hidden_slides_export_preview_and_thumbnail(tmp_path, monkeypatch):
    monkeypatch.setattr(powerpoint_export.sys, "platform", "linux")
    source = tmp_path / "original.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(600), Pt(300)
    for text in ("Visible", "Hidden"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Pt(20), Pt(20), Pt(200), Pt(50)).text = text
    deck.slides[1]._element.set("show", "0")
    deck.save(source)
    before = source.read_bytes()
    calls = []

    def convert(working, destination, output_format):
        calls.append(output_format)
        assert Presentation(working).slides[1]._element.get("show") == "1"
        target = destination / "preview.pdf"
        pdf = canvas.Canvas(str(target), pagesize=(600, 300))
        for text in ("Visible", "Hidden"):
            pdf.drawString(20, 250, text)
            pdf.showPage()
        pdf.save()
        return target

    monkeypatch.setattr(powerpoint_export, "convert_office", convert)
    preview, notice = service.office_preview(source, tmp_path)
    assert preview and "LibreOffice" in notice
    assert source.read_bytes() == before
    assert "ExportHiddenSlides" in calls[0]
    manifest = json.loads((tmp_path / "powerpoint-export.json").read_text())
    assert manifest["renderer"] == "LibreOffice Impress"
    (tmp_path / "slides.json").write_text(
        json.dumps(
            {
                "slides": [
                    {
                        "number": n,
                        "preview": f"editable-preview/page-{n}.html",
                        "width": 600,
                        "height": 300,
                    }
                    for n in (1, 2)
                ]
            }
        )
    )
    monkeypatch.setattr(thumbnails, "renderer", lambda: pytest.fail("Linux must not invoke Swift"))
    for number in (1, 2):
        target = thumbnails.thumbnail(tmp_path, number)
        assert target.read_bytes().startswith(b"RIFF")
        assert target.stat().st_size <= 2000
        assert thumbnails.thumbnail(tmp_path, number) == target
    assert len(calls) == 1


@pytest.mark.parametrize("failure", ["missing", "timeout", "no_output"])
def test_linux_office_failure_is_reported(tmp_path, monkeypatch, failure):
    monkeypatch.setattr(
        office_linux.shutil,
        "which",
        lambda name: None if failure == "missing" else "/usr/bin/soffice",
    )

    def run(*args, **kwargs):
        if failure == "timeout":
            raise subprocess.TimeoutExpired(args[0], 150)
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(office_linux.subprocess, "run", run)
    with pytest.raises(RuntimeError):
        office_linux.convert_office(tmp_path / "original.docx", tmp_path / "out", "html")


def test_https_origin_requires_explicit_configuration(tmp_path, monkeypatch):
    monkeypatch.setattr(desk_config, "DATA", tmp_path)
    client = TestClient(web.app, base_url="https://localhost")
    assert (
        client.post(
            "/api/folders", json={"name": "HTTPS"}, headers={"Origin": "https://localhost"}
        ).status_code
        == 403
    )
    monkeypatch.setattr(desk_config, "ALLOWED_ORIGINS", {"https://localhost"})
    assert (
        client.post(
            "/api/folders", json={"name": "HTTPS"}, headers={"Origin": "https://localhost"}
        ).status_code
        == 201
    )
    assert (
        client.post(
            "/api/folders", json={"name": "Other"}, headers={"Origin": "https://attacker.example"}
        ).status_code
        == 403
    )
    assert client.get("/health/live").json() == {"status": "ok"}
    monkeypatch.setattr(desk_config, "MODELS", tmp_path / "missing-models")
    assert client.get("/health/ready").status_code == 503


def test_calc_hidden_columns_and_rows_do_not_expand_merged_cells(tmp_path, monkeypatch):
    source = tmp_path / "original.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Visible"
    sheet.append(["Heading"])
    sheet.merge_cells("A1:C1")
    sheet.append(["Hidden"])
    sheet.append(["English", "Secret", "日本語", "=1+1"])
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True
    workbook.create_sheet("Hidden").sheet_state = "hidden"
    workbook.save(source)
    before = source.read_bytes()

    def convert(source, destination, output_format):
        target = destination / "original.html"
        target.write_text(
            '<html><head><meta charset="utf-8"></head><body><table>'
            '<colgroup span="3" width="70"></colgroup>'
            '<tr><td colspan="3">Heading</td><td></td></tr>'
            "<tr><td>English</td><td>日本語</td><td>2</td></tr>"
            "</table></body></html>"
        )
        return target

    monkeypatch.setattr(office_linux, "convert_office", convert)
    preview = office_linux.office_html(source, tmp_path)
    sheets = preview_sheets(tmp_path, preview)
    assert len(sheets) == 1
    tree = html.fromstring(sheet_html(tmp_path, "a" * 32, sheets[0]))
    assert tree.xpath('//tr[@class="sheet-column-axis"]/th/text()') == ["A", "C", "D"]
    assert tree.xpath('//th[@class="sheet-row-axis"]/text()') == ["1", "3"]
    assert tree.xpath('string(//td[@colspan="2"])') == "Heading"
    assert source.read_bytes() == before


def test_translation_keeps_source_row_after_a_merge_spans_hidden_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Heading"
    sheet.merge_cells("A1:A3")
    sheet["C1"] = "Top"
    sheet["C3"] = "Bottom"
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True
    original = tmp_path / "original.xlsx"
    workbook.save(original)
    workbook = load_workbook(original)
    raw = (
        '<html><body><table class="worksheet">'
        '<tr><th class="sheet-row-axis">1</th><td rowspan="2">Heading</td><td>Top</td></tr>'
        '<tr><th class="sheet-row-axis">3</th><td>Bottom</td></tr>'
        "</table></body></html>"
    )
    segments, excluded = worksheet_segments(raw, workbook, 1)
    assert [(segment["cell_ref"], segment["source_text"]) for segment in segments] == [
        ("A1", "Heading"),
        ("C1", "Top"),
        ("C3", "Bottom"),
    ]
    assert excluded == 0
