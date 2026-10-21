import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
import reportlab
from lxml import html
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont as ReportlabFont
from reportlab.pdfgen import canvas

from docling_desk.documents.conversion import office_preview
from docling_desk.preview import pdfium_access
from docling_desk.preview.editable_preview import (
    NOTICE,
    RENDERER_VERSION,
    build_editable_preview,
    editable_preview,
    ensure_pdf_preview,
    refresh_saved_preview,
    source_textboxes,
)
from docling_desk.preview.native_preview import build_native_preview
from docling_desk.preview.slides import split_quicklook
from docling_desk.translation.source import html_segments, slide_template
from docling_desk.translation.view import apply_text


def sources(folder, pdf_count=2, size=(960, 540)):
    source, pdf = folder / "original.pptx", folder / "powerpoint-rendered.pdf"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(960), Pt(540)
    labels = ["Material", "Vendor & <review>"]
    for label in labels:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Pt(100), Pt(100), Pt(200), Pt(80))
        shape.text = label
    deck.save(source)
    font_path = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    pdfmetrics.registerFont(ReportlabFont("EditablePreviewVera", str(font_path)))
    document = canvas.Canvas(str(pdf), pagesize=size)
    for label in labels[:pdf_count]:
        document.setFillColorRGB(0.5, 0.5, 0.9)
        document.rect(100, 360, 200, 80, fill=1, stroke=1)
        document.setFillColorRGB(0, 0, 0)
        document.setFont("EditablePreviewVera", 14)
        document.drawString(107, 420, label)
        document.showPage()
    document.save()
    return source, pdf


def test_preview_is_a_page_image_with_an_invisible_selectable_text_layer(tmp_path, monkeypatch):
    source, pdf = sources(tmp_path)
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, pdf)]
    preview = build_editable_preview(source, pdf, tmp_path)
    assert editable_preview(source, tmp_path) == preview
    frames = split_quicklook(tmp_path / preview, tmp_path, 2)
    for number, text in [(1, "Material"), (2, "Vendor & <review>")]:
        tree = html.fromstring((tmp_path / frames[number]).read_text())
        # The graphics live in the image; the only DOM text is the invisible layer.
        assert tree.xpath("//img/@src") == [f"page-{number}.webp"]
        assert tree.xpath('//div[contains(@class,"pdf-text")]/text()') == [text]
        assert tree.xpath('//div[contains(@class,"pdf-text")]/@data-source-id') == ["shape-2"]
        assert not tree.xpath("//svg|//text|//tspan|//script|//review|//iframe|//link")
        assert "@font-face" not in (tmp_path / frames[number]).read_text()
        with Image.open(tmp_path / f"editable-preview/page-{number}.webp") as image:
            assert image.format == "WEBP" and image.size == (1920, 1080)
            # The blue rectangle drawn in the PDF is part of the image.
            assert image.convert("RGB").getpixel((250, 330)) != (255, 255, 255)
    manifest = json.loads((tmp_path / "editable-preview/manifest.json").read_text())
    assert manifest["version"] == RENDERER_VERSION == 4
    assert manifest["pages"] == 2 and manifest["text_blocks"] == manifest["mapped_blocks"] == 2
    assert {"page-1.html", "page-1.webp", "page-2.webp"} <= set(manifest["assets"])
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, pdf)]

    def no_quicklook(*args, **kwargs):
        raise AssertionError("Editable native rendering must take priority")

    monkeypatch.setattr("docling_desk.documents.conversion.subprocess.run", no_quicklook)
    assert office_preview(source, tmp_path) == (preview, NOTICE)


def test_text_block_sits_at_the_powerpoint_text_frame_in_page_percentages(tmp_path):
    source, pdf = sources(tmp_path)
    preview = build_editable_preview(source, pdf, tmp_path)
    tree = html.fromstring(
        (tmp_path / split_quicklook(tmp_path / preview, tmp_path, 2)[1]).read_text()
    )
    block = tree.xpath('//div[contains(@class,"pdf-text")]')[0]
    style = dict(item.split(":", 1) for item in block.get("style").split(";"))
    assert float(style["left"].rstrip("%")) == pytest.approx(100 / 960 * 100, abs=0.01)
    assert float(style["top"].rstrip("%")) == pytest.approx(100 / 540 * 100, abs=0.01)
    assert float(style["width"].rstrip("%")) == pytest.approx(200 / 960 * 100, abs=0.01)
    assert float(style["height"].rstrip("%")) == pytest.approx(80 / 540 * 100, abs=0.01)
    # Font size follows the rendered text (14pt) and scales with the page width.
    assert float(style["font-size"].removesuffix("cqw")) == pytest.approx(14 / 960 * 100, abs=0.01)
    assert block.get("data-size") == "14.000"


def test_translation_replaces_the_block_text_and_draws_it_over_the_image(tmp_path):
    source, pdf = sources(tmp_path)
    preview = build_editable_preview(source, pdf, tmp_path)
    frames = split_quicklook(tmp_path / preview, tmp_path, 2)
    raw = slide_template(tmp_path, frames[2])
    segments = html_segments(raw, html.fromstring(raw).xpath('//div[@class="slide"]'))
    assert len(segments) == 1
    assert segments[0]["source_text"] == "Vendor & <review>"
    assert segments[0]["locator"]["slot"] == "text"
    translated = apply_text(raw, {"segments": segments}, {segments[0]["id"]: "仕入先"})
    tree, original = html.fromstring(translated), html.fromstring(raw)
    block = tree.xpath('//div[contains(@class,"pdf-text")]')[0]
    assert block.text == "仕入先" and block.get("data-translated") == "true"
    assert "data-translated" not in original.xpath('//div[contains(@class,"pdf-text")]')[0].attrib
    assert block.get("style").startswith(
        original.xpath('//div[contains(@class,"pdf-text")]/@style')[0]
    )
    # The page image is untouched, and so is every other attribute of the block.
    assert tree.xpath("//img/@src") == original.xpath("//img/@src")
    assert block.get("data-source-id") == "shape-2"
    stale = dict(segments[0], source_text="different source")
    with pytest.raises(ValueError):
        apply_text(raw, {"segments": [stale]}, {stale["id"]: "仕入先"})


def test_long_translation_shrinks_to_fit_the_frame_but_not_below_a_floor(tmp_path):
    source, pdf = sources(tmp_path)
    preview = build_editable_preview(source, pdf, tmp_path)
    raw = slide_template(tmp_path, split_quicklook(tmp_path / preview, tmp_path, 2)[1])
    segments = html_segments(raw, html.fromstring(raw).xpath('//div[@class="slide"]'))

    def size(text):
        tree = html.fromstring(apply_text(raw, {"segments": segments}, {segments[0]["id"]: text}))
        block = tree.xpath('//div[contains(@class,"pdf-text")]')[0]
        return float(block.get("style").rsplit("font-size:", 1)[1].removesuffix("cqw"))

    base = 14 / 960 * 100
    assert size("品目") == pytest.approx(base, abs=1e-3)
    assert size("これは枠に収まるよう文字を小さくする長い訳文です。" * 3) < base
    assert size("あ" * 400) == pytest.approx(base * 0.55, abs=1e-3)


def test_pdf_text_outside_every_powerpoint_frame_becomes_its_own_block(tmp_path):
    source, pdf = sources(tmp_path)
    document = canvas.Canvas(str(pdf), pagesize=(960, 540))
    document.setFont("Helvetica", 14)
    document.drawString(107, 420, "Material")
    document.drawString(40, 30, "Master footer")
    document.showPage()
    document.drawString(107, 420, "Vendor & <review>")
    document.showPage()
    document.save()
    preview = build_editable_preview(source, pdf, tmp_path)
    frames = split_quicklook(tmp_path / preview, tmp_path, 2)
    tree = html.fromstring((tmp_path / frames[1]).read_text())
    blocks = tree.xpath('//div[contains(@class,"pdf-text")]')
    assert {b.text: b.get("data-source-id") for b in blocks} == {
        "Material": "shape-2",
        "Master footer": None,
    }


def test_pdf_line_inside_an_owned_frame_does_not_become_a_second_block(tmp_path):
    source, pdf = sources(tmp_path)
    document = canvas.Canvas(str(pdf), pagesize=(960, 540))
    document.setFont("Helvetica", 14)
    document.drawString(107, 420, "Material")
    document.drawString(107, 380, "decoded differently")  # inside the frame, not in its text
    document.showPage()
    document.drawString(107, 420, "Vendor & <review>")
    document.showPage()
    document.save()
    preview = build_editable_preview(source, pdf, tmp_path)
    frames = split_quicklook(tmp_path / preview, tmp_path, 2)
    tree = html.fromstring((tmp_path / frames[1]).read_text())
    blocks = tree.xpath('//div[contains(@class,"pdf-text")]')
    assert [b.text for b in blocks] == ["Material"]
    assert blocks[0].get("data-source-id") == "shape-2"


def test_group_offsets_and_scaling_are_retained_for_translation_targets():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Pt(10), Pt(20), Pt(200), Pt(80))
    child.text = "Nested label"
    group.left, group.top, group.width, group.height = Pt(100), Pt(200), Pt(400), Pt(160)
    box = source_textboxes(slide)[0]
    assert box["shape_path"] == [group.shape_id, child.shape_id]
    assert box["bbox"] == pytest.approx([100, 200, 500, 360])
    assert box["frame"] == pytest.approx([100, 200, 400, 160])
    assert box["angle"] == pytest.approx(0)


def test_rotated_textbox_keeps_its_own_size_and_angle():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Pt(100), Pt(100), Pt(200), Pt(40))
    shape.text = "Turned"
    shape.rotation = 90
    box = source_textboxes(slide)[0]
    # The bounding box swaps, the frame keeps the shape's own width and height.
    assert box["bbox"] == pytest.approx([180, 20, 220, 220])
    assert box["frame"][2:] == pytest.approx([200, 40])
    assert box["angle"] == pytest.approx(90)
    assert box["frame"][0] + box["frame"][2] / 2 == pytest.approx(200)


@pytest.mark.parametrize("pdf_count,size", [(1, (960, 540)), (2, (540, 960))])
def test_wrong_pdf_rejected_before_registration(tmp_path, pdf_count, size):
    source, pdf = sources(tmp_path, pdf_count, size)
    with pytest.raises(ValueError):
        build_editable_preview(source, pdf, tmp_path)
    assert not (tmp_path / "editable-preview").exists()


def test_encrypted_and_over_limit_pdfs_are_refused(tmp_path, monkeypatch):
    locked = tmp_path / "locked.pdf"
    page = canvas.Canvas(str(locked), encrypt=StandardEncryption("user", "owner"))
    page.drawString(50, 50, "secret")
    page.save()
    with pytest.raises(ValueError, match="暗号化"):
        ensure_pdf_preview(locked, tmp_path)
    assert not (tmp_path / "editable-preview").exists()
    source, pdf = sources(tmp_path)
    monkeypatch.setattr(pdfium_access, "MAX_PAGES", 1)
    with pytest.raises(ValueError, match="上限"):
        build_editable_preview(source, pdf, tmp_path)
    assert not (tmp_path / "editable-preview").exists()
    assert not list(tmp_path.glob(".editable-preview-*"))


@pytest.mark.parametrize("change", ["source", "pdf", "asset", "outside", "image"])
def test_stale_or_incomplete_editable_cache_is_rejected(tmp_path, change):
    source, pdf = sources(tmp_path)
    build_editable_preview(source, pdf, tmp_path)
    if change in {"source", "pdf"}:
        target = source if change == "source" else pdf
        target.write_bytes(target.read_bytes() + b"changed")
    elif change in {"asset", "image"}:
        (tmp_path / f"editable-preview/page-2.{'html' if change == 'asset' else 'webp'}").unlink()
    else:
        path = tmp_path / "editable-preview/manifest.json"
        data = json.loads(path.read_text())
        data["assets"] = ["../outside.svg"]
        path.write_text(json.dumps(data))
    assert editable_preview(source, tmp_path) is None


def test_preview_of_an_older_renderer_is_never_read_and_is_rebuilt_from_the_saved_pdf(tmp_path):
    source, pdf = sources(tmp_path)
    preview = build_editable_preview(source, pdf, tmp_path)
    directory = tmp_path / "editable-preview"
    # What an older renderer left behind: its own SVG pages and a lower version.
    (directory / "page-1.html").write_text(
        "<html><body><svg><text>old svg</text></svg></body></html>"
    )
    manifest = json.loads((directory / "manifest.json").read_text())
    manifest["version"] = 3
    (directory / "manifest.json").write_text(json.dumps(manifest))
    assert editable_preview(source, tmp_path) is None
    assert split_quicklook(tmp_path / preview, tmp_path, 2) == {}
    with pytest.raises(ValueError, match="再生成"):
        slide_template(tmp_path, "editable-preview/page-1.html")
    # Without a verified saved export nothing is rebuilt and the old page stays unread.
    saved = pdf.read_bytes()
    pdf.write_bytes(saved + b"tampered")
    (tmp_path / "powerpoint-export.json").write_text(
        json.dumps({"source_sha256": "x", "pdf_sha256": "y"})
    )
    assert refresh_saved_preview(source, tmp_path) is None
    pdf.write_bytes(saved)
    (tmp_path / "powerpoint-export.json").unlink()
    manifest["pdf"] = "powerpoint-rendered.pdf"
    (directory / "manifest.json").write_text(json.dumps(manifest))
    assert refresh_saved_preview(source, tmp_path) == preview
    assert "old svg" not in (directory / "page-1.html").read_text()
    assert json.loads((directory / "manifest.json").read_text())["version"] == RENDERER_VERSION
    assert slide_template(tmp_path, "editable-preview/page-1.html")


def test_reference_image_cache_is_never_the_default_preview(tmp_path, monkeypatch):
    source, pdf = sources(tmp_path)
    reference = build_native_preview(source, pdf, tmp_path)
    monkeypatch.setattr(
        "docling_desk.documents.conversion.subprocess.run",
        lambda *a, **k: SimpleNamespace(returncode=1),
    )
    assert office_preview(source, tmp_path)[0] != reference
