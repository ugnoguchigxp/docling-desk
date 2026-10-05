import hashlib
from copy import deepcopy

from lxml import html
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from reportlab.pdfgen import canvas

from docling_desk.preview.quicklook import corrected_preview, prune_repeated_triangles


def make_preview(tmp_path, text="Body text"):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(20), Pt(30), Pt(200), Pt(60))
    shape.text_frame.margin_left = shape.text_frame.margin_right = Pt(7)
    shape.text_frame.margin_top = shape.text_frame.margin_bottom = Pt(3)
    shape.text_frame.text = "\nBody text\n"
    paragraph = shape.text_frame.paragraphs[1]
    paragraph.runs[0].font.size = Pt(14)
    paragraph._p.get_or_add_pPr().set("indent", "-279400")  # 22pt hanging indent.
    source = tmp_path / "original.pptx"
    deck.save(source)
    preview = tmp_path / "Preview.html"
    preview.write_text(
        "<html><head><style>.inner.inner.inner{display:table;width:200;height:60;"
        "margin-left:7px;margin-right:7px;margin-top:3px;margin-bottom:3px}"
        ".body.body.body{font-size:24;text-indent:-22px;margin-left:22px;}"
        '.run.run.run{font-size:14;font-family:"Segoe UI";}'
        ".bullet.bullet.bullet{margin-left:-22px;padding-right:38px;}</style></head>"
        '<body><div class="slide"><div style="position:absolute;left:20;top:30;'
        'width:200;height:60"><div class="inner"><div><p class="run"> </p>'
        f'<p class="body"><span class="bullet">•</span><span class="run">{text}</span></p>'
        '<p class="run"> </p></div></div></div></div></body></html>',
        encoding="utf-8",
    )
    return source, preview


def test_corrects_matching_source_without_changing_text_or_originals(tmp_path):
    source, preview = make_preview(tmp_path)
    original = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, preview)]
    corrected = corrected_preview(source, preview)
    tree = html.fromstring(corrected.read_text())
    assert "width:186.000px !important" in tree.xpath('//div[@class="inner"]/@style')[0]
    body = tree.xpath('//p[@class="body"]')[0]
    assert "font-size:14px !important" in body.get("style")
    assert "line-height:1.2 !important" in body.get("style")
    assert "margin-left:0px !important" in body[0].get("style")
    assert "width:22px !important" in body[0].get("style")
    assert all(
        "line-height:0 !important" in p.get("style") for p in tree.xpath('//p[@class="run"]')
    )
    assert tree.text_content() == html.fromstring(preview.read_text()).text_content()
    assert original == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, preview)]
    before = corrected.read_bytes()
    assert corrected_preview(source, preview).read_bytes() == before


def test_different_text_does_not_receive_source_paragraph_spacing(tmp_path):
    source, preview = make_preview(tmp_path, text="Different text")
    corrected = corrected_preview(source, preview)
    tree = html.fromstring(corrected.read_text())
    assert "Different text" in tree.text_content()
    assert all("line-height:0" not in p.get("style", "") for p in tree.xpath("//p"))


def test_page_count_mismatch_keeps_original_preview(tmp_path):
    source, preview = make_preview(tmp_path)
    preview.write_text(preview.read_text().replace('<div class="slide">', "<div>"))
    assert corrected_preview(source, preview) == preview
    assert not preview.with_name("Preview-corrected.html").exists()


def test_pdf_diagram_keeps_transparency_and_rejects_outside_assets(tmp_path):
    local = tmp_path / "local"
    local.mkdir()
    source, preview = make_preview(local)
    pdf = canvas.Canvas(str(local / "arrow.pdf"), pagesize=(100, 50))
    pdf.setFillColorRGB(0.5, 0.5, 0.8)
    pdf.rect(10, 10, 80, 20, fill=1, stroke=0)
    pdf.save()
    (tmp_path / "outside.pdf").write_bytes((local / "arrow.pdf").read_bytes())
    preview.write_text(
        preview.read_text().replace(
            "</body>",
            '<img src="arrow.pdf"><img src="../outside.pdf"><img src="https://example.test/arrow.pdf"></body>',
        )
    )
    tree = html.fromstring(corrected_preview(source, preview).read_text())
    assert tree.xpath("//img/@src") == [
        "arrow.quicklook.png",
        "../outside.pdf",
        "https://example.test/arrow.pdf",
    ]
    assert not (tmp_path / "outside.quicklook.png").exists()
    with Image.open(local / "arrow.quicklook.png") as image:
        assert image.mode == "RGBA"
        assert image.getextrema()[-1] == (0, 255)


def test_invalid_pdf_is_reported_without_removing_diagram(tmp_path):
    source, preview = make_preview(tmp_path)
    (tmp_path / "broken.pdf").write_bytes(b"invalid")
    preview.write_text(preview.read_text().replace("</body>", '<img src="broken.pdf"></body>'))
    tree = html.fromstring(corrected_preview(source, preview).read_text())
    assert tree.xpath("//img/@src") == ["broken.pdf"]
    assert '"asset": "broken.pdf"' in preview.with_name("preview-corrections.json").read_text()


def test_reused_vector_belongs_to_source_page_and_unknown_vectors_are_retained():
    deck = Presentation()
    for _ in range(3):
        deck.slides.add_slide(deck.slide_layouts[6])
    # The real owner is later than the erroneous occurrence in HTML.
    deck.slides[2].shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Pt(40), Pt(50), Pt(100), Pt(20))
    slides = [
        html.fromstring(
            '<div class="slide"><p>Text remains</p>'
            '<img src="arrow.pdf" style="left:20;top:30;width:140;height:60">'
            '<img src="unknown.pdf" style="left:100;top:100;width:50;height:50"></div>'
        )
        for _ in range(3)
    ]
    removed = prune_repeated_triangles(slides, deck)
    assert removed == [
        {"slide": 1, "asset": "arrow.pdf", "source_slides": [3]},
        {"slide": 2, "asset": "arrow.pdf", "source_slides": [3]},
    ]
    assert slides[2].xpath(".//img/@src") == ["arrow.pdf", "unknown.pdf"]
    assert all(s.xpath('.//img[@src="unknown.pdf"]') for s in slides)
    assert all(s.text_content() == "Text remains" for s in slides)


def test_inherited_and_group_triangles_prevent_speculative_asset_removal():
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    second = deck.slides.add_slide(deck.slide_layouts[6])
    first.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Pt(40), Pt(50), Pt(100), Pt(20))
    group = second.shapes.add_group_shape()
    group.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Pt(0), Pt(0), Pt(10), Pt(10))
    slides = [
        html.fromstring(
            '<div class="slide"><img src="arrow.pdf" style="left:20;top:30;width:140;height:60"></div>'
        )
        for _ in range(2)
    ]
    assert prune_repeated_triangles(slides, deck) == []
    # A layout triangle is also inherited by both slides.
    second._element.spTree.remove(group._element)
    first.slide_layout._element.spTree.append(deepcopy(first.shapes[0]._element))
    assert prune_repeated_triangles(slides, deck) == []
