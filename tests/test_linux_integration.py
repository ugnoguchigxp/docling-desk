"""Exercise the real Linux Office renderer in the Docker test environment."""

import hashlib
import json
import shutil
import sys

import pytest
from lxml import html
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Pt

from docling_desk.documents.conversion import Job, convert_job
from docling_desk.preview.sheets import preview_sheets, sheet_html
from docling_desk.preview.thumbnails import thumbnail
from docling_desk.translation.source import source_map


@pytest.mark.skipif(
    sys.platform != "linux" or shutil.which("libreoffice") is None,
    reason="Requires the Linux container's LibreOffice renderer",
)
def test_real_office_hidden_content_merges_and_translation_binding(tmp_path):
    folder = tmp_path / "edge-sheet"
    folder.mkdir()
    source = folder / "original.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Visible"
    sheet.append(["Heading"])
    sheet.merge_cells("A1:C1")
    sheet.append(["Hidden"])
    sheet.append(["English", "Secret", "日本語", "=1+1"])
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["B"].hidden = True
    workbook.create_sheet("Hidden").sheet_state = "hidden"
    workbook.create_sheet("Last").append(["Last"])
    workbook.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    job = Job(id="a" * 32, filename="edge.xlsx")
    convert_job(folder, job)
    assert job.state == "success", job
    sheets = preview_sheets(folder, job.preview)
    assert [sheet["number"] for sheet in sheets] == [1, 3]
    tree = html.fromstring(sheet_html(folder, job.id, sheets[0]))
    assert tree.xpath('//tr[@class="sheet-column-axis"]/th/text()') == ["A", "C", "D"]
    assert tree.xpath('//th[@class="sheet-row-axis"]/text()') == ["1", "3"]
    assert tree.xpath('string(//td[@colspan="2"])') == "Heading"
    segments = source_map(folder)["units"][0]["segments"]
    assert any(
        segment["cell_ref"] == "C3" and segment["source_text"] == "日本語" for segment in segments
    )
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before

    folder = tmp_path / "edge-slide"
    folder.mkdir()
    source = folder / "original.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(600), Pt(300)
    for text in ("Visible", "Hidden"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Pt(30), Pt(30), Pt(200), Pt(60)).text = text
    deck.slides[1]._element.set("show", "0")
    deck.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    job = Job(id="b" * 32, filename="edge.pptx")
    convert_job(folder, job)
    assert job.state == "success", job
    assert json.loads((folder / "editable-preview/manifest.json").read_text())["pages"] == 2
    assert thumbnail(folder, 2).is_file()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
