import hashlib
import json
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from lxml import html

import docling_desk.app as web
import docling_desk.documents.conversion as service
from docling_desk import config as desk_config
from docling_desk.documents.conversion import Job, save_job, validate_input
from docling_desk.documents.tables import project_tables
from docling_desk.explanation.source import snapshot
from docling_desk.translation.source import source_map
from docling_desk.translation.store import atomic_json, result_path


def make_docx(path):
    doc = Document()
    doc.add_heading("Word copy test", 0)
    doc.add_paragraph("日本語の本文。Sales increased by 120 units.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "区分", "件数"
    cells = table.add_row().cells
    cells[0].text, cells[1].text = "A", "120"
    doc.add_page_break()
    doc.add_heading("Second section", 1)
    doc.add_paragraph("次の本文。")
    doc.save(path)


@pytest.fixture
def word(tmp_path, monkeypatch):
    folder = tmp_path / ("d" * 32)
    folder.mkdir()
    make_docx(folder / "original.docx")

    # Test extraction without depending on the operating system's renderer.
    def preview(source, directory):
        target = directory / "quicklook/Preview.html"
        target.parent.mkdir()
        target.write_text(
            '<html><head><meta charset="utf-8"><style>p{color:blue}</style></head>'
            '<body onload="alert(1)"><p>Word copy test</p><p>日本語の本文。</p>'
            "<table><tr><td>A</td><td>120</td></tr></table><p>次の本文。</p>"
            '<script>alert(1)</script><iframe src="https://example.test"></iframe></body></html>'
        )
        return target.relative_to(directory).as_posix(), "Test preview"

    monkeypatch.setattr(service, "office_preview", preview)
    job = Job(id=folder.name, filename="sample.docx")
    service.convert_job(folder, job)
    assert job.state == "success", job.error
    return folder, job


def test_docx_validation_rejects_mislabeled_office_and_corrupt_files(tmp_path):
    source = tmp_path / "original.docx"
    make_docx(source)
    validate_input(source)
    with ZipFile(source, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(ValueError, match="Office構造"):
        validate_input(source)
    source.write_bytes(b"invalid")
    with pytest.raises(ValueError, match="ZIP"):
        validate_input(source)


def test_word_extraction_tables_rag_and_document_unit(word):
    folder, job = word
    assert job.pages == 0  # Do not invent print page provenance.
    assert job.tables == 1 and not job.slide_layout
    assert "Sales increased by 120" in (folder / "text.txt").read_text()
    assert "次の本文" in (folder / "text.txt").read_text()
    chunks = [json.loads(line) for line in (folder / "rag.jsonl").read_text().splitlines()]
    assert any("120" in chunk["text"] for chunk in chunks)
    from docling_core.types.doc import DoclingDocument

    table = project_tables(DoclingDocument.load_from_json(folder / "document.json"), job.filename)[
        0
    ]
    assert table.source_label == "文書全体" and table.pages == []
    assert table.rows == [["区分", "件数"], ["A", "120"]]
    source = source_map(folder)
    assert source["unlocated_count"] == 0
    assert len(source["units"]) == 1
    unit = source["units"][0]
    assert (unit["id"], unit["kind"], unit["mode"]) == ("document-1", "document", "replace")
    assert "次の本文。" in [s["source_text"] for s in unit["segments"]]
    explanation = snapshot(folder)
    assert explanation["unlocated_count"] == 0
    texts = [b["text"] for b in explanation["units"][0]["blocks"]]
    assert any("Sales increased" in text for text in texts)
    assert any("次の本文" in text for text in texts)
    assert len(explanation["units"][0]["tables"]) == 1


def test_word_routes_copy_bridge_translation_and_original_preservation(word, monkeypatch):
    folder, job = word
    monkeypatch.setattr(desk_config, "DATA", folder.parent)
    original = hashlib.sha256((folder / "original.docx").read_bytes()).hexdigest()
    client = TestClient(web.app)
    response = client.get(f"/view/{job.id}/word")
    assert response.status_code == 200
    tree = html.fromstring(response.text)
    assert tree.xpath("//body/@data-unit") == ["document-1"]
    assert tree.xpath("//script/@src") == ["/static/document-frame.js", "/static/file-drop.js"]
    assert not tree.xpath("//iframe|//@onload")
    assert "allow-same-origin" not in response.headers["content-security-policy"]
    for name in ["document-frame.js", "document-frame.css"]:
        assert client.get("/static/" + name).status_code == 200
    unit = source_map(folder)["units"][0]
    record = {
        "unit_id": unit["id"],
        "state": "completed",
        "result": {
            "source_hash": unit["source_hash"],
            "translations": {s["id"]: "Translated " + s["source_text"] for s in unit["segments"]},
        },
    }
    atomic_json(result_path(folder, "en", "document-1"), record)
    translated = client.get(f"/view/{job.id}/word?language=en")
    assert translated.headers["x-translation-state"] == "translated"
    assert "Translated 日本語の本文。" in translated.text
    assert client.get(f"/view/{job.id}/word?language=invalid").status_code == 422
    assert (
        client.get(f"/files/{job.id}/original.docx?download=true").content
        == (folder / "original.docx").read_bytes()
    )
    assert original == hashlib.sha256((folder / "original.docx").read_bytes()).hexdigest()
    job.filename = "renamed.docx"
    job.original_filename = "sample.docx"
    save_job(folder, job)
    assert client.get(f"/view/{job.id}/word").status_code == 200


def test_word_translation_uses_extracted_text_when_quicklook_is_unavailable(word, monkeypatch):
    folder, job = word
    job.preview = None
    job.state = "partial"
    save_job(folder, job)
    monkeypatch.setattr(desk_config, "DATA", folder.parent)
    source = source_map(folder)
    unit = source["units"][0]
    assert unit["id"] == "document-1" and unit["mode"] == "panel"
    texts = [segment["source_text"] for segment in unit["segments"]]
    assert "次の本文。" in texts
    assert any("Sales increased by 120" in text for text in texts)
    assert "区分" in texts and "A" in texts
    client = TestClient(web.app)
    assert client.get(f"/api/jobs/{job.id}/translations").status_code == 200
    atomic_json(
        result_path(folder, "en", "document-1"),
        {
            "unit_id": "document-1",
            "state": "completed",
            "result": {
                "source_hash": unit["source_hash"],
                "translations": {
                    segment["id"]: "Translated " + segment["source_text"]
                    for segment in unit["segments"]
                },
            },
        },
    )
    translated = client.get(f"/api/jobs/{job.id}/translations/en/document-1").json()
    assert translated["available"] and translated["mode"] == "panel"
    assert "Translated 次の本文。" in translated["texts"]
    assert client.get(f"/view/{job.id}/word").status_code == 404


def test_upload_accepts_docx_and_keeps_folder_placement(tmp_path, monkeypatch):
    from types import SimpleNamespace

    from docling_desk.api import uploads

    monkeypatch.setattr(desk_config, "DATA", tmp_path / "data")
    desk_config.DATA.mkdir()
    # This test checks upload placement; capacity boundaries have their own tests.
    usage = uploads.shutil.disk_usage(desk_config.DATA)
    monkeypatch.setattr(uploads.shutil, "disk_usage", lambda _: usage._replace(free=3 * 1024**3))
    submissions = []
    source = tmp_path / "sample.docx"
    make_docx(source)
    with TestClient(web.app) as client:
        client.app.state.worker = SimpleNamespace(submit=lambda *args: submissions.append(args))
        parent = client.post("/api/folders", json={"name": "Word"}).json()["id"]
        response = client.post(
            "/api/upload",
            data={"folder_id": parent},
            files={"file": ("sample.DOCX", source.read_bytes())},
        )
        assert response.status_code == 202, response.text
        job = response.json()
        assert job["folder_id"] == parent and len(submissions) == 1
        assert (
            desk_config.DATA / "content" / "documents" / job["id"] / "original.docx"
        ).read_bytes() == source.read_bytes()
        if submissions:
            client.app.state.slots.release()
