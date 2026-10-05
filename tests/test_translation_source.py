import json

from lxml import html
from openpyxl import Workbook

from docling_desk.preview.sheets import preview_sheets
from docling_desk.translation.source import html_segments, source_map, template_for


def test_source_maps_are_stable_and_unit_boundaries_are_preserved(translation_document):
    for kind in ("slide", "sheet", "page"):
        folder = translation_document(kind)
        first = source_map(folder)
        assert first == source_map(folder)
        assert [u["id"] for u in first["units"]] == [f"{kind}-1", f"{kind}-2"]
        assert all(u["segments"] for u in first["units"])
        for unit in first["units"]:
            ids = [s["id"] for s in unit["segments"]]
            assert len(ids) == len(set(ids))


def test_duplicate_text_and_nested_runs_have_distinct_bindings():
    raw = "<html><body><p><b>同じ</b>文字</p><p><i>同じ</i></p></body></html>"
    roots = html.fromstring(raw).xpath("//p")
    segments = html_segments(raw, roots)
    assert [s["source_text"] for s in segments] == ["同じ", "文字", "同じ"]
    assert len({(s["locator"]["node_path"], s["locator"]["slot"]) for s in segments}) == 3


def test_stylesheet_changes_invalidate_existing_binding(translation_document):
    folder = translation_document()
    first = source_map(folder)
    slides = json.loads((folder / "slides.json").read_text())["slides"]
    stylesheet = (folder / slides[0]["preview"]).parent / "slide-layout.css"
    stylesheet.write_text(stylesheet.read_text() + "\np{color:red}")
    assert first["units"][0]["source_hash"] != source_map(folder)["units"][0]["source_hash"]


def test_excel_protects_numbers_formula_hidden_columns_and_merged_cells(translation_document):
    folder = translation_document("sheet")
    workbook = Workbook()
    ws = workbook.active
    ws.title = "集計"
    ws["A1"] = "結合文字"
    ws.merge_cells("A1:C1")
    ws["A2"] = "非表示"
    ws.row_dimensions[2].hidden = True
    ws["A3"] = "=1+1"
    ws["B3"] = "非表示列"
    ws.column_dimensions["B"].hidden = True
    ws["C3"] = "文字"
    ws["A4"] = 123
    ws["C4"] = "確認"
    workbook.create_sheet("注記")
    workbook.save(folder / "original.xlsx")
    job = json.loads((folder / "job.json").read_text())
    sheets = preview_sheets(folder, job["preview"])
    sheets[0]["path"].write_text(
        '<html><body><table class="worksheet"><col><col><tr><td colspan="2"><b>結合</b>文字</td></tr><tr><td>2</td><td>文字</td></tr><tr><td>123</td><td>確認</td></tr></table></body></html>'
    )
    source = source_map(folder)
    segments = source["units"][0]["segments"]
    assert [s["source_text"] for s in segments] == ["結合", "文字", "文字", "確認"]
    assert [s["cell_ref"] for s in segments] == ["A1", "A1", "C3", "C4"]
    assert source["units"][0]["excluded_count"] == 0


def test_uncertain_cells_are_excluded_instead_of_guessed(translation_document):
    folder = translation_document("sheet")
    job = json.loads((folder / "job.json").read_text())
    sheet = preview_sheets(folder, job["preview"])[0]
    sheet["path"].write_text(
        '<html><body><table class="worksheet"><tr><td>一致しない本文</td></tr></table></body></html>'
    )
    unit = source_map(folder)["units"][0]
    assert unit["segments"] == []
    assert unit["excluded_count"] == 1


def test_office_template_keeps_quirks_mode(translation_document):
    folder = translation_document()
    job = json.loads((folder / "job.json").read_text())
    raw = template_for(folder, job, 1)
    assert not raw.lower().startswith("<!doctype")


def test_image_slide_uses_existing_page_text_as_panel(translation_document):
    folder = translation_document()
    slides = json.loads((folder / "slides.json").read_text())["slides"]
    path = folder / slides[0]["preview"]
    path.write_text('<html><body><img src="Attachment1.png"></body></html>')
    unit = source_map(folder)["units"][0]
    assert unit["mode"] == "panel" and unit["segments"]
    assert all("node_path" not in s["locator"] for s in unit["segments"])
