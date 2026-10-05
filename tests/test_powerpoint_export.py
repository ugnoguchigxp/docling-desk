import hashlib
import json
import subprocess
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Pt
from reportlab.pdfgen import canvas

import docling_desk.preview.powerpoint_export as export
from docling_desk.documents.conversion import office_preview
from docling_desk.preview.editable_preview import build_editable_preview


def source(folder):
    path = folder / "original.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(600), Pt(300)
    deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_textbox(
        Pt(30), Pt(30), Pt(200), Pt(60)
    ).text = "First source"
    deck.save(path)
    return path


def write_pdf(path, pages=1):
    doc = canvas.Canvas(str(path), pagesize=(600, 300))
    for _ in range(pages):
        doc.drawString(30, 250, "First source")
        doc.showPage()
    doc.save()


@pytest.fixture
def native_renderer(tmp_path, monkeypatch):
    app = tmp_path / "PowerPoint.app"
    app.mkdir()
    workspace = tmp_path / "sandbox"
    monkeypatch.setattr(export, "POWERPOINT", app)
    monkeypatch.setattr(export, "WORKSPACE", workspace)
    monkeypatch.setattr(export.sys, "platform", "darwin")
    calls = []

    def run(command, **kwargs):
        calls.append(command)
        assert Path(command[2]).parent.is_relative_to(workspace)
        assert Path(command[2]).name != "original.pptx"
        assert command[4] == Path(command[2]).name
        assert kwargs["timeout"] == 150
        write_pdf(Path(command[3]))
        return SimpleNamespace(returncode=0, stderr="")

    monkeypatch.setattr(export.subprocess, "run", run)
    return calls


def test_new_pptx_automatically_uses_native_pdf_and_text_html(tmp_path, native_renderer):
    path = source(tmp_path)
    before = path.read_bytes()
    preview, notice = office_preview(path, tmp_path)
    assert preview == "editable-preview/Preview.html"
    assert "テキスト" in notice
    assert "First source" in (tmp_path / preview).read_text()
    assert len(native_renderer) == 1
    assert path.read_bytes() == before
    assert not list(export.WORKSPACE.iterdir())
    assert not (tmp_path / "quicklook").exists()
    assert office_preview(path, tmp_path)[0] == preview
    assert len(native_renderer) == 1
    deck = Presentation(path)
    deck.slides[0].shapes[0].text = "Changed source"
    deck.save(path)
    assert office_preview(path, tmp_path)[0] == preview
    assert len(native_renderer) == 2
    manifest = json.loads((tmp_path / "powerpoint-export.json").read_text())
    assert manifest["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_verified_older_pdf_can_rebuild_html_without_rerender(tmp_path, monkeypatch):
    path = source(tmp_path)
    pdf = tmp_path / "powerpoint-rendered.pdf"
    write_pdf(pdf)
    build_editable_preview(path, pdf, tmp_path)
    manifest = tmp_path / "editable-preview/manifest.json"
    data = json.loads(manifest.read_text())
    data["version"] = 1
    manifest.write_text(json.dumps(data))
    monkeypatch.setattr(
        export.subprocess, "run", lambda *a, **k: pytest.fail("Valid PDF must be reused")
    )
    assert office_preview(path, tmp_path)[0] == "editable-preview/Preview.html"
    assert json.loads(manifest.read_text())["version"] == 2


@pytest.mark.parametrize("failure", ["missing", "timeout", "wrong_pages"])
def test_failed_native_export_is_visible_and_never_falls_back_to_quicklook(
    tmp_path, native_renderer, monkeypatch, failure
):
    path = source(tmp_path)
    old = tmp_path / "powerpoint-rendered.pdf"
    old.write_bytes(b"previous export")
    if failure == "missing":
        monkeypatch.setattr(export, "POWERPOINT", tmp_path / "unavailable")
    else:

        def run(command, **kwargs):
            if failure == "timeout":
                raise subprocess.TimeoutExpired(command, 150)
            write_pdf(Path(command[3]), pages=2)
            return SimpleNamespace(returncode=0, stderr="")

        monkeypatch.setattr(export.subprocess, "run", run)
    preview, notice = office_preview(path, tmp_path)
    assert preview is None and "作成できません" in notice
    assert not (tmp_path / "quicklook").exists()
    assert old.read_bytes() == b"previous export"


def test_native_script_only_closes_its_named_copy():
    script = export.SCRIPT.read_text()
    assert "presentation workingName" in script
    assert "active presentation" not in script
    assert "quit" not in script


def test_hidden_slide_visibility_changes_only_in_the_disposable_copy(tmp_path):
    original = source(tmp_path)
    deck = Presentation(original)
    deck.slides[0]._element.set("show", "0")
    deck.save(original)
    before = original.read_bytes()
    copy = tmp_path / "copy.pptx"
    export.prepare_copy(original, copy)
    assert Presentation(copy).slides[0]._element.get("show") == "1"
    with ZipFile(original) as a, ZipFile(copy) as b:
        assert a.namelist() == b.namelist()
        assert all(
            a.read(name) == b.read(name) for name in a.namelist() if name != "ppt/slides/slide1.xml"
        )
    assert original.read_bytes() == before
