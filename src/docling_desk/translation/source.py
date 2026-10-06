"""Bind visible text to its existing HTML node; never recreate document layout."""

from __future__ import annotations

import json
import re
from functools import lru_cache
from pathlib import Path

from docling_core.types.doc import DoclingDocument, TableItem, TextItem
from lxml import html
from openpyxl import load_workbook

from docling_desk.documents.text import DOCUMENT_SUFFIXES
from docling_desk.preview.editable_preview import RENDERER_VERSION, ensure_pdf_preview
from docling_desk.preview.hashing import digest
from docling_desk.preview.sheets import preview_sheets, sheet_html
from docling_desk.storage import job_file, original_file
from docling_desk.translation.store import LOCK, atomic_json, fingerprint

BINDING_VERSION = 2


def text_candidate(value: str) -> bool:
    return any(c.isalpha() for c in value)


def serialize(tree, raw: str) -> str:
    # libxml supplies a default HTML4 doctype even when the source has none.
    # Keep the source's rendering mode rather than emitting that default.
    declaration = re.match(r"\s*(<!doctype[^>]*>)", raw, re.I)
    return html.tostring(tree, encoding="unicode", doctype=declaration[1] if declaration else None)


def slide_template(folder: Path, relative: str) -> str:
    if not relative:
        raise ValueError("原本プレビューがありません。")
    folder = folder.resolve()
    path = (folder / relative).resolve()
    if not path.is_relative_to(folder.resolve()):
        raise ValueError("スライドのプレビューが不正です。")
    raw = path.read_text(encoding="utf-8")
    tree = html.fromstring(raw)
    for node in tree.xpath("//script|//iframe|//object|//embed|//base|//form|//meta[@http-equiv]"):
        node.drop_tree()
    for node in tree.iter():
        for attribute in list(node.attrib):
            if attribute.lower().startswith("on"):
                del node.attrib[attribute]
        for attribute in ("src", "href"):
            value = node.get(attribute)
            if value is None or value.startswith(("data:image/", "#")):
                continue
            target = (path.parent / value).resolve()
            if target.is_relative_to(folder.resolve()) and target.is_file():
                node.set(attribute, f"/files/{folder.name}/{target.relative_to(folder).as_posix()}")
            else:
                del node.attrib[attribute]
    return serialize(tree, raw)


@lru_cache(maxsize=8)
def read_slide_previews(path: str, stat: tuple) -> dict:
    # Keep only the page-to-template index, not all extracted slide blocks.
    return {
        s["number"]: {"preview": s.get("preview")}
        for s in json.loads(Path(path).read_text())["slides"]
    }


def slide_previews(folder: Path) -> dict:
    path = folder / "slides.json"
    return read_slide_previews(str(path), signature(path))


def template_for(folder: Path, job: dict, number: int) -> str:
    suffix = Path(job.get("original_filename") or job["filename"]).suffix.lower()
    if suffix == ".pptx":
        slide = slide_previews(folder)[number]
        if not slide.get("preview"):
            raise ValueError("原本プレビューがありません。抽出本文を別枠で翻訳します。")
        return slide_template(folder, slide["preview"])
    if suffix == ".pdf":
        ensure_pdf_preview(original_file(folder, ".pdf"), folder)
        return slide_template(folder, f"editable-preview/page-{number}.html")
    if suffix in DOCUMENT_SUFFIXES:
        if number != 1:
            raise ValueError("この形式は文書全体を表示します。")
        return slide_template(folder, job["preview"])
    sheets = preview_sheets(folder, job["preview"])
    sheet = next(s for s in sheets if s["number"] == number)
    return sheet_html(folder, folder.name, sheet)


def html_segments(raw: str, roots: list) -> list[dict]:
    tree = html.fromstring(raw)
    scope = roots or [tree.find("body")]
    segments = []
    # Callers use roots parsed from this exact string, with paths into this tree.
    for root in scope:
        for node in root.iter():
            if not isinstance(node.tag, str) or node.tag in {"script", "style"}:
                continue
            if any(p.get("data-glyph-clusters") == "true" for p in node.iterancestors()):
                continue
            if node.get("data-glyph-clusters") == "true":
                value = "".join(node.itertext())
                if text_candidate(value):
                    segments.append(
                        {
                            "id": f"t{len(segments) + 1:05d}",
                            "source_text": value,
                            "locator": {
                                "node_path": node.getroottree().getpath(node),
                                "slot": "svg_text",
                            },
                        }
                    )
                continue
            for slot in ("text", "tail"):
                value = getattr(node, slot)
                if node is root and slot == "tail":
                    continue
                if value and text_candidate(value):
                    segments.append(
                        {
                            "id": f"t{len(segments) + 1:05d}",
                            "source_text": value,
                            "locator": {
                                "node_path": node.getroottree().getpath(node),
                                "slot": slot,
                            },
                        }
                    )
    return segments


def worksheet_segments(raw: str, workbook, number: int) -> tuple[list[dict], int]:
    tree = html.fromstring(raw)
    worksheet = workbook.worksheets[number - 1]
    segments, excluded = [], 0
    hidden_columns = set()
    for dimension in worksheet.column_dimensions.values():
        if dimension.hidden:
            hidden_columns.update(range(dimension.min, dimension.max + 1))
    visible_columns = [
        n for n in range(1, max(worksheet.max_column, 1) + 1) if n not in hidden_columns
    ]
    occupied = set()
    for table in tree.xpath('//table[contains(concat(" ", @class, " "), " worksheet ")]'):
        rows = table.xpath("./tr|./tbody/tr")
        visible_rows = [row for row in rows if row.xpath('./th[@class="sheet-row-axis"]')]
        row_numbers = [
            int(row.xpath('string(./th[@class="sheet-row-axis"])')) for row in visible_rows
        ]
        for row_index, row in enumerate(visible_rows):
            axis = row.xpath('./th[@class="sheet-row-axis"]')
            if not axis:
                continue
            row_no = int(axis[0].text)
            col_index = 0
            for cell in row.xpath("./td"):
                while (row_no, col_index) in occupied:
                    col_index += 1
                colspan, rowspan = int(cell.get("colspan", "1")), int(cell.get("rowspan", "1"))
                for r in row_numbers[row_index : row_index + rowspan]:
                    for c in range(col_index, col_index + colspan):
                        occupied.add((r, c))
                candidate = None
                if col_index < len(visible_columns):
                    candidate = worksheet.cell(row_no, visible_columns[col_index])
                value = "".join(cell.itertext()).strip()
                # Require exact source text as well as position and type. Uncertain cells are never guessed.
                if (
                    candidate is not None
                    and candidate.data_type in {"s", "inlineStr"}
                    and isinstance(candidate.value, str)
                    and value == candidate.value.strip()
                ):
                    bound = html_segments(raw, [cell])
                    for segment in bound:
                        segment["id"] = f"t{len(segments) + 1:05d}"
                        segment["cell_ref"] = candidate.coordinate
                        segments.append(segment)
                elif text_candidate(value) and (
                    candidate is None or candidate.data_type in {"s", "inlineStr"}
                ):
                    excluded += 1
                col_index += colspan
    return segments, excluded


def document_segments(doc: DoclingDocument, number: int | None) -> list[dict]:
    segments = []
    for item, _ in doc.iterate_items():
        if not isinstance(item, (TextItem, TableItem)):
            continue
        if number is not None and (not item.prov or item.prov[0].page_no != number):
            continue
        parent = item.parent.resolve(doc) if item.parent else None
        if parent is not None and parent.label.value == "picture":
            continue
        values = (
            [(item.text, {"ref": item.self_ref})]
            if isinstance(item, TextItem)
            else [
                (
                    cell.text,
                    {
                        "ref": item.self_ref,
                        "row": cell.start_row_offset_idx,
                        "col": cell.start_col_offset_idx,
                    },
                )
                for cell in item.data.table_cells
            ]
        )
        for value, locator in values:
            if value and text_candidate(value):
                segments.append(
                    {"id": f"t{len(segments) + 1:05d}", "source_text": value, "locator": locator}
                )
    return segments


def build_source(folder: Path) -> dict:
    job = json.loads((job_file(folder)).read_text())
    suffix = Path(job.get("original_filename") or job["filename"]).suffix.lower()
    original = original_file(folder, suffix)
    kind = (
        "document"
        if suffix in DOCUMENT_SUFFIXES
        else {".pptx": "slide", ".xlsx": "sheet", ".pdf": "page"}[suffix]
    )
    doc = DoclingDocument.load_from_json(folder / "document.json")
    units = []
    workbook = load_workbook(original, data_only=False) if kind == "sheet" else None
    numbers = sorted(doc.pages)
    slides = {}
    if kind == "document":
        numbers = [1]
    if kind == "slide":
        slides = slide_previews(folder)
        numbers = list(slides)
    if kind == "sheet":
        numbers = [s["number"] for s in preview_sheets(folder, job["preview"])]
    try:
        for number in numbers:
            unit = {
                "id": f"{kind}-{number}",
                "kind": kind,
                "number": number,
                "mode": "panel",
                "excluded_count": 0,
            }
            raw = ""
            preview_available = (
                bool(slides[number].get("preview"))
                if kind == "slide"
                else bool(job.get("preview")) or kind == "page"
            )
            if not preview_available:
                unit["preview_unavailable_reason"] = (
                    "原本プレビューがありません。抽出本文を別枠で翻訳します。原本上の文字置換は利用できません。"
                )
            if preview_available:
                try:
                    if kind == "slide":
                        raw = slide_template(folder, slides[number]["preview"])
                    elif kind == "page":
                        raw = slide_template(folder, f"editable-preview/page-{number}.html")
                    else:
                        raw = template_for(folder, job, number)
                except (OSError, ValueError):
                    unit["preview_unavailable_reason"] = (
                        "原本プレビューを読み込めません。抽出本文を別枠で翻訳します。原本上の文字置換は利用できません。"
                    )
                raw = raw or "<html><body></body></html>"
                tree = html.fromstring(raw)
                if kind == "sheet":
                    segments, excluded = worksheet_segments(raw, workbook, number)
                    unit.update(segments=segments, excluded_count=excluded, mode="replace")
                else:
                    roots = tree.xpath(
                        '//div[contains(concat(" ",normalize-space(@class)," ")," slide ")]'
                    )
                    segments = html_segments(raw, roots)
                    if segments:
                        unit.update(segments=segments, mode="replace")
                assets = {}
                for url in tree.xpath('//link[@rel="stylesheet"]/@href'):
                    prefix = f"/files/{folder.name}/"
                    if url.startswith(prefix):
                        path = (folder / url[len(prefix) :]).resolve()
                        if path.is_relative_to(folder.resolve()) and path.is_file():
                            assets[url] = cached_digest(path)
                unit["template_hash"] = fingerprint([raw, assets])
            if unit["mode"] == "panel":
                unit["segments"] = document_segments(doc, None if kind == "document" else number)
            if not unit.get("segments"):
                unit["unavailable_reason"] = "このページには翻訳できる抽出本文がありません。"
            unit["source_hash"] = fingerprint([BINDING_VERSION, unit])
            units.append(unit)
    finally:
        if workbook:
            workbook.close()
    return {
        "schema_version": 1,
        "binding_version": BINDING_VERSION,
        "document_id": folder.name,
        "original_sha256": digest(original),
        "unlocated_count": sum(
            isinstance(i, (TextItem, TableItem)) and not i.prov for i, _ in doc.iterate_items()
        )
        if kind != "document"
        else 0,
        "units": units,
    }


@lru_cache(maxsize=8192)
def file_digest(path: str, signature: tuple) -> str:
    return digest(Path(path))


def signature(path: Path) -> tuple:
    stat = path.stat()
    return stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns, stat.st_ino


def cached_digest(path: Path) -> str:
    return file_digest(str(path), signature(path))


@lru_cache(maxsize=4)
def read_source(path: str, stat: tuple) -> dict:
    return json.loads(Path(path).read_text())


def dependencies(folder: Path) -> dict:
    """Portable content hashes; local stats only avoid rereading unchanged bytes.

    Preview directories include CSS, fonts, images and other reference assets.
    Inventory changes also invalidate the binding. Runtime translation results,
    job state, search data and the binding itself do not participate.
    """
    job = json.loads(job_file(folder).read_text())
    suffix = Path(job.get("original_filename") or job["filename"]).suffix.lower()
    paths = {"original": original_file(folder, suffix), "document.json": folder / "document.json"}
    for name in ("slides.json", "powerpoint-rendered.pdf", "powerpoint-export.json"):
        if (folder / name).is_file():
            paths[name] = folder / name
    preview_dirs = {"editable-preview", "quicklook", "office-preview"}
    if job.get("preview"):
        preview_dirs.add(Path(job["preview"]).parts[0])
    # Older slide templates can live beside the other derived files.
    extensions = {
        ".html",
        ".htm",
        ".svg",
        ".css",
        ".js",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".gif",
        ".ttf",
        ".otf",
        ".woff",
        ".woff2",
        ".bmp",
        ".avif",
        ".pdf",
    }
    for path in folder.rglob("*"):
        if not path.is_file() or not path.resolve().is_relative_to(folder):
            continue
        relative = path.relative_to(folder)
        if relative.parts[0] in {"translations", "explanations", "thumbnails"} or any(
            part.endswith((".tmp", ".part", ".lock")) for part in relative.parts
        ):
            continue
        if relative.parts[0] in preview_dirs or path.suffix.lower() in extensions:
            paths[relative.as_posix()] = path
    return {
        "binding_version": BINDING_VERSION,
        "renderer_version": RENDERER_VERSION,
        "cache_version": 2,
        "job": {k: job.get(k) for k in ("filename", "original_filename", "preview")},
        "files": {name: cached_digest(path) for name, path in sorted(paths.items())},
    }


def source_map(folder: Path) -> dict:
    folder = folder.resolve()
    with LOCK:
        path = folder / "translation-source.json"
        current = dependencies(folder)
        if path.exists():
            try:
                saved = read_source(str(path), signature(path))
                if saved.get("dependencies") == current and saved.get("document_id") == folder.name:
                    return saved
            except (OSError, ValueError, TypeError):
                pass
        job = json.loads(job_file(folder).read_text())
        if Path(job.get("original_filename") or job["filename"]).suffix.lower() == ".pdf":
            # Generate / upgrade once before capturing template dependencies,
            # rather than checking or generating the complete PDF for every page.
            ensure_pdf_preview(original_file(folder, ".pdf"), folder)
            current = dependencies(folder)
        source = build_source(folder)
        # A source edit invalidates every stored result, even if one unit's visible text happens to match.
        # PDF preview creation may have added dependencies during build_source.
        after = dependencies(folder)
        # Fail rather than publishing bindings to a source edited mid-build.
        if any(after["files"].get(k) != v for k, v in current["files"].items()):
            raise ValueError("展開中に原本または参照素材が変更されました。再度開いてください。")
        dependencies_hash = fingerprint(after)
        for unit in source["units"]:
            unit["source_hash"] = fingerprint([unit["source_hash"], dependencies_hash])
        source["dependencies"] = after
        atomic_json(path, source)
        return source
