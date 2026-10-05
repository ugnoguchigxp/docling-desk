"""Build the Word counterpart of the demo's synthetic PDF/PPTX/XLSX data."""

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def make_word_sample(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in [("Normal", 11), ("Title", 24), ("Heading 1", 16), ("Heading 2", 12)]:
        style = document.styles[name]
        style.font.name = "Hiragino Sans W3"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Hiragino Sans W3")
        for color in style.element.xpath(".//w:color"):
            for attr in ["themeColor", "themeTint", "themeShade"]:
                color.attrib.pop(qn("w:" + attr), None)
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = 1.25
    normal.space_after = Pt(8)
    document.add_paragraph("Wordの表示とコピーを確認する", "Title")
    document.add_paragraph(
        "本文・表・図を含むWordの合成サンプルです。PDF・PowerPoint・Excelのサンプルと"
        "同じ架空データを使っています。文字選択、コピー、表の操作、翻訳・解説を確認できます。"
    )
    document.add_heading("区分ごとの件数", 1)
    document.add_paragraph("合計は450件です。Bは180件で最も多く、Aより60件多くなっています。")
    document.add_paragraph("表1 区分ごとの件数")
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(10.6), Cm(6)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "区分", "件数"
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for label, count in [("A", "120"), ("B", "180"), ("C", "150"), ("合計", "450")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, count
    for i, row in enumerate(table.rows):
        for col, cell in enumerate(row.cells):
            cell.width = Cm(10.6 if col == 0 else 6)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders, margins = OxmlElement("w:tcBorders"), OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                border = OxmlElement("w:" + edge)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn("w:" + key), value)
                borders.append(border)
                margin = OxmlElement("w:" + edge)
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.extend([borders, margins])
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DCEAF0" if i == 0 else "FFFFFF")
            props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                )
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.bold = i == 0 or i == 4
    document.add_paragraph()
    picture = document.add_picture(str(output.parent / "chart.png"), width=Cm(13.8))
    picture._inline.docPr.set("descr", "区分Aは120件、Bは180件、Cは150件の棒グラフ")
    document.add_paragraph("図1 表1と同じ件数を示す棒グラフ")
    document.add_page_break()
    document.add_heading("確認する操作", 1)
    document.add_paragraph(
        "これは印刷時の2ページ目に置いた本文です。アプリではWordを文書全体として表示するため、"
        "ページ番号を付けて分割しません。後半の本文も抽出・翻訳・解説に含まれることを確認してください。"
    )
    for text in [
        "本文の一部を選び、上部のコピーボタンか⌘C／Ctrl+Cでコピーする。",
        "「表を操作」でA・B・Cの件数と合計450件を確認し、範囲選択してコピーする。",
        "「−」「＋」「100%」「幅に合わせる」で表示倍率を変える。",
        "「構造・参照」「RAGデータ」を開き、この後半の本文も含まれることを確認する。",
    ]:
        document.add_paragraph(text, "List Bullet")
    document.add_heading("英語の本文", 1)
    document.add_paragraph(
        "Category B has 180 records. The total is 450 records. These figures are fictional."
    )
    document.add_heading("確認時の注意", 1)
    document.add_paragraph(
        "画像の棒グラフに含まれる文字は直接選択できません。数値は表からコピーできます。"
        "翻訳・解説を新しく生成する場合は、本文と表を画面に表示されたAIサービスへ送信します。"
    )
    document.core_properties.title = "Wordの表示とコピーを確認する"
    document.core_properties.author = "Docling Desk"
    document.core_properties.created = datetime(2026, 10, 3, tzinfo=timezone.utc)
    document.core_properties.modified = document.core_properties.created
    document.save(str(output))


if __name__ == "__main__":
    make_word_sample(Path.cwd() / "samples/synthetic-report.docx")
