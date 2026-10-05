"""Generate tiny deterministic software-test fixtures, never real customer material."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen.canvas import Canvas

from docling_desk.development.word import make_word_sample

OUT = Path.cwd() / "samples"
OUT.mkdir(exist_ok=True)
chart = Image.new("RGB", (600, 270), "white")
draw = ImageDraw.Draw(chart)
for i, (label, amount) in enumerate([("A", 120), ("B", 180), ("C", 150)]):
    x = 75 + 170 * i
    draw.rectangle((x, 230 - amount, x + 90, 230), fill="#216d87")
    draw.text((x + 35, 240), label, fill="black")
    draw.text((x + 30, 212 - amount), str(amount), fill="black")
chart.save(OUT / "chart.png")
pdfmetrics.registerFont(UnicodeCIDFont("HeiseiMin-W3"))
canvas = Canvas(str(OUT / "synthetic-report.pdf"), pagesize=(595, 842))
canvas.setTitle("合成サンプル: 日本語・表・図")
canvas.setFont("HeiseiMin-W3", 22)
canvas.drawString(45, 785, "合成サンプル: 月別の件数")
canvas.setFont("HeiseiMin-W3", 12)
canvas.drawString(45, 749, "これは抽出機能を確認するための架空の資料です。")
canvas.drawString(45, 716, "表1\u3000区分ごとの件数。合計は450件です。")
rows = [("区分", "件数"), ("A", "120"), ("B", "180"), ("C", "150"), ("合計", "450")]
for n, row in enumerate(rows):
    y = 670 - n * 35
    for col, text in enumerate(row):
        canvas.rect(45 + col * 190, y - 10, 190, 35, fill=0)
        canvas.drawString(58 + col * 190, y, text)
canvas.drawImage(ImageReader(chart), 45, 275, width=480, height=216)
canvas.drawString(45, 248, "図1\u3000表1の件数を棒グラフで示しました。")
canvas.drawString(45, 213, "Bの件数が最も多く、Aより60件多いことが分かります。")
canvas.setFont("HeiseiMin-W3", 10)
canvas.drawString(45, 40, "合成テスト資料 / 1")
canvas.showPage()
canvas.setFont("HeiseiMin-W3", 22)
canvas.drawString(45, 785, "図と表の参照を確認する")
canvas.setFont("HeiseiMin-W3", 12)
canvas.drawString(45, 740, "前ページの表1と図1は同じ架空データを示しています。")
canvas.drawString(45, 710, "日本語の抽出、読み順、ページ番号、表構造を確認してください。")
canvas.drawString(45, 680, "意味関係が自動で完全に保持されるとは限りません。")
canvas.setFont("HeiseiMin-W3", 10)
canvas.drawString(45, 40, "合成テスト資料 / 2")
canvas.save()
prs = Presentation()
prs.slide_width = Inches(12)
prs.slide_height = Inches(6.75)
for i in range(2):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(11), Inches(0.7))
    box.text_frame.text = ["合成サンプル: 表と図の比較", "日本語と出典の確認"][i]
    for paragraph in box.text_frame.paragraphs:
        paragraph.font.size = Pt(28)
        paragraph.font.name = "Hiragino Sans"
    if i == 0:
        table = slide.shapes.add_table(5, 2, Inches(0.6), Inches(1.5), Inches(4), Inches(3)).table
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
        slide.shapes.add_picture(str(OUT / "chart.png"), Inches(5.2), Inches(1.8), width=Inches(6))
        body = "図1: 表1の架空データを比較。BはAより60件多い。"
    else:
        body = "これは合成テスト資料です。\n表1と図1の対応、日本語、ページ番号を確認します。\n翻訳と実資料の品質検証は未実施です。"
    box = slide.shapes.add_textbox(
        Inches(0.6), Inches(5.2 if i == 0 else 1.7), Inches(11), Inches(1)
    )
    box.text_frame.text = body
    for p in box.text_frame.paragraphs:
        p.font.size = Pt(20)
        p.font.name = "Hiragino Sans"
prs.save(str(OUT / "synthetic-slides.pptx"))
wb = Workbook()
ws = wb.active
ws.title = "集計"
ws.append(["合成サンプル: 区分ごとの件数"])
ws.append(["架空データ / 図1は表1と同じ数値"])
ws.append([])
ws.append(["区分", "件数"])
for row in rows[1:]:
    ws.append([row[0], int(row[1])])
ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 18
for cell in ws[4]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="216D87")
ws.freeze_panes = "A5"
ws.add_image(ExcelImage(str(OUT / "chart.png")), "D4")
notes = wb.create_sheet("注記")
notes.append(["合成テスト資料"])
notes.append(["表1と図1は同じ架空データです。"])
notes.append(["実資料の品質検証は未実施。"])
notes.column_dimensions["A"].width = 60
wb.save(OUT / "synthetic-sheet.xlsx")
make_word_sample(OUT / "synthetic-report.docx")
print("Generated PDF, PPTX, XLSX and DOCX synthetic samples in", OUT)
