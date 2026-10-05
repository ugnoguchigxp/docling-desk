"""Isolated headless LibreOffice exports for Linux document previews."""

from __future__ import annotations

import html
import shutil
import subprocess
from pathlib import Path
from tempfile import TemporaryDirectory

from lxml import html as markup
from openpyxl import load_workbook

from docling_desk.preview.sheets import workbook_sheets


def normalize_merged_cells(table, worksheet, sheet: dict) -> None:
    """Calc keeps source spans even when its HTML omits hidden rows/columns."""
    columns = [
        number
        for number in range(1, worksheet.max_column + 1)
        if number not in sheet["hidden_columns"]
    ]
    merged = {(area.min_row, area.min_col): area for area in worksheet.merged_cells.ranges}
    occupied = set()
    row_number = 0
    for row in table.xpath("./tr|./tbody/tr"):
        row_number += 1
        while row_number in sheet["hidden_rows"]:
            row_number += 1
        column_index = 0
        for cell in row.xpath("./td"):
            while (row_number, column_index) in occupied:
                column_index += 1
            area = (
                merged.get((row_number, columns[column_index]))
                if column_index < len(columns)
                else None
            )
            if area is not None:
                span = sum(area.min_col <= number <= area.max_col for number in columns)
                rows = [
                    number
                    for number in range(area.min_row, area.max_row + 1)
                    if number not in sheet["hidden_rows"]
                ]
                cell.set("colspan", str(span))
                cell.set("rowspan", str(len(rows)))
                for number in rows:
                    for index in range(column_index, column_index + span):
                        occupied.add((number, index))
            column_index += int(cell.get("colspan", "1"))


def convert_office(source: Path, destination: Path, output_format: str) -> Path:
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        raise RuntimeError("LibreOfficeがありません。Linux用Dockerイメージを使用してください。")
    destination.mkdir(parents=True, exist_ok=True)
    # Never reuse a user profile or attach to another running Office instance.
    with TemporaryDirectory(prefix="docling-office-") as work:
        profile = Path(work) / "profile"
        try:
            result = subprocess.run(
                [
                    binary,
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nofirststartwizard",
                    "--convert-to",
                    output_format,
                    "--outdir",
                    str(destination.resolve()),
                    str(source.resolve()),
                ],
                capture_output=True,
                text=True,
                timeout=150,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("LibreOfficeの変換が時間内に完了しませんでした。") from exc
    output = destination / f"{source.stem}.{output_format.split(':', 1)[0]}"
    if result.returncode != 0 or not output.is_file() or not output.stat().st_size:
        raise RuntimeError("LibreOfficeの変換結果を確認できませんでした。")
    return output


def office_html(source: Path, folder: Path) -> str:
    directory = folder / "libreoffice"
    # Export into a fresh directory: a failed conversion cannot reuse stale HTML.
    with TemporaryDirectory(prefix="office-html-", dir=folder) as work:
        raw = convert_office(source, Path(work), "html")
        if source.suffix.lower() == ".xlsx":
            tree = markup.document_fromstring(raw.read_bytes())
            tables = tree.xpath("//body/table")
            visible = [sheet for sheet in workbook_sheets(source) if sheet["visible"]]
            workbook = load_workbook(source)
            if len(tables) != len(visible):
                raise ValueError("LibreOfficeのシート数が原本と一致しません。")
            sections = []
            for sheet, table in zip(visible, tables, strict=True):
                table.set("class", "worksheet")
                normalize_merged_cells(table, workbook.worksheets[sheet["number"] - 1], sheet)
                target = Path(work) / f"sheet-{sheet['number']}.html"
                head = tree.find("head")
                target.write_text(
                    "<!doctype html><html>"
                    + (
                        markup.tostring(head, encoding="unicode")
                        if head is not None
                        else "<head></head>"
                    )
                    + "<body>"
                    + markup.tostring(table, encoding="unicode")
                    + "</body></html>",
                    encoding="utf-8",
                )
                sections.append(
                    f"<section><h2>{html.escape(sheet['name'])}</h2>"
                    f'<iframe src="{target.name}" sandbox></iframe></section>'
                )
            raw.write_text(
                '<!doctype html><html><head><meta charset="utf-8"></head><body>'
                + "".join(sections)
                + "</body></html>",
                encoding="utf-8",
            )
        if directory.exists():
            shutil.rmtree(directory)
        shutil.copytree(work, directory)
    return (directory / raw.name).relative_to(folder).as_posix()
