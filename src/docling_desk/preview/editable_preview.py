"""Show PDF and PowerPoint pages as a page image with an invisible, selectable text layer.

Every page is rendered once by PDFium into a WebP image. A positioned
``div.pdf-text`` per text block sits on top of it, transparent, so the text
stays selectable and copyable. Translations replace a block's text and make the
block opaque (see ``translation.view.apply_text``); nothing here reconstructs
vector text or fonts.
"""

from __future__ import annotations

import html
import json
import math
import re
import shutil
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from statistics import median
from tempfile import TemporaryDirectory
from threading import Lock
from uuid import uuid4

from docling.utils.locks import pypdfium2_lock
from PIL import Image
from pptx import Presentation
from pptx.shapes.group import GroupShape

from docling_desk.preview.hashing import digest
from docling_desk.preview.pdf_text import TextLine, group_blocks, read_page
from docling_desk.preview.pdfium_access import checked_pdf, open_pdf
from docling_desk.storage import original_file

NOTICE = "原本側はPowerPointの描画をページ画像として表示しています。文字は画像の上に選択できる文字層を重ねており、翻訳すると文字枠の位置に訳文を重ねて表示します。"
PDF_NOTICE = "PDFの描画をページ画像として表示しています。文字は画像の上に選択できる文字層を重ねており、翻訳すると文字の位置に訳文を重ねて表示します。スキャン画像内の文字は原本画像のままです。"
RENDERER_VERSION = 4
PDF_RENDER_LOCK = Lock()
IMAGE_SCALE = 2.0
IMAGE_MAX_SIDE = 3200
IMAGE_MAX_PIXELS = 12_000_000
IMAGE_QUALITY = 85
# Translation backgrounds below this share of one colour are drawn translucent.
UNIFORM_SHARE = 0.45
# A text area is sampled at about this many pixels, however large it is drawn.
SAMPLE_PIXELS = 2500
FONT_STACK = 'Arial,"Hiragino Sans","Yu Gothic","Noto Sans CJK JP",sans-serif'
PAGE_STYLE = (
    "html,body{margin:0;width:100%;height:100%;overflow:hidden;background:#fff}"
    ".page-canvas{position:absolute;inset:0;container-type:inline-size;overflow:hidden;background:#fff}"
    ".page-canvas>img{position:absolute;inset:0;width:100%;height:100%;display:block;"
    "-webkit-user-select:none;user-select:none;pointer-events:none}"
    ".pdf-text{position:absolute;margin:0;box-sizing:border-box;display:flex;"
    "flex-direction:column;color:transparent;white-space:pre;"
    f"font-family:{FONT_STACK};"
    "-webkit-user-select:text;user-select:text}"
    ".pdf-text[data-translated]{color:var(--fg);background:var(--bg);background-clip:content-box;"
    "white-space:pre-wrap;overflow-wrap:anywhere}"
)


def editable_preview(source: Path, folder: Path) -> str | None:
    """The page-preview index when a complete current-version cache matches the exact bytes."""
    directory = folder / "editable-preview"
    try:
        manifest = json.loads((directory / "manifest.json").read_text())
        pdf = (
            original_file(folder, ".pdf")
            if manifest["pdf"] == "original.pdf"
            else folder / manifest["pdf"]
        ).resolve()
        if (
            not pdf.is_relative_to(folder.resolve())
            and pdf != original_file(folder, ".pdf").resolve()
        ) or manifest["version"] != RENDERER_VERSION:
            return None
        if digest(source) != manifest["source_sha256"] or digest(pdf) != manifest["pdf_sha256"]:
            return None
        assets = manifest["assets"]
        if not assets or not all(
            (directory / name).resolve().is_relative_to(directory.resolve())
            and (directory / name).is_file()
            for name in assets
        ):
            return None
        return (directory / "Preview.html").relative_to(folder).as_posix()
    except (OSError, ValueError, KeyError, TypeError):
        return None


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", text).casefold()


def source_textboxes(slide) -> list[dict]:
    """Record source IDs and absolute boxes, including scaled/rotated groups and table cells."""
    boxes = []

    def visit(shapes, matrix, path):
        for shape in shapes:
            a, b, c, d, e, f = matrix
            x, y, w, h = [
                float(value) for value in (shape.left, shape.top, shape.width, shape.height)
            ]
            current_path = [*path, shape.shape_id]
            if isinstance(shape, GroupShape):
                transform = shape._element.grpSpPr.xfrm
                if transform.chExt is None or not transform.chExt.cx or not transform.chExt.cy:
                    continue
                sx, sy = transform.cx / transform.chExt.cx, transform.cy / transform.chExt.cy
                angle = math.radians(shape.rotation)
                ca, sa = math.cos(angle), math.sin(angle)
                ga, gb, gc, gd = ca * sx, sa * sx, -sa * sy, ca * sy
                ge = (
                    x
                    + w / 2
                    - ca * w / 2
                    + sa * h / 2
                    - ga * transform.chOff.x
                    - gc * transform.chOff.y
                )
                gf = (
                    y
                    + h / 2
                    - sa * w / 2
                    - ca * h / 2
                    - gb * transform.chOff.x
                    - gd * transform.chOff.y
                )
                visit(
                    shape.shapes,
                    (
                        a * ga + c * gb,
                        b * ga + d * gb,
                        a * gc + c * gd,
                        b * gc + d * gd,
                        a * ge + c * gf + e,
                        b * ge + d * gf + f,
                    ),
                    current_path,
                )
                continue
            frames = []
            if shape.has_text_frame:
                frames.append((shape.text_frame, x, y, w, h, None))
            if shape.has_table:
                table = shape.table
                top = y
                for row_index, row in enumerate(table.rows):
                    left = x
                    for column_index, column in enumerate(table.columns):
                        cell = table.cell(row_index, column_index)
                        if not cell.is_spanned:
                            cw = sum(
                                table.columns[i].width
                                for i in range(column_index, column_index + cell.span_width)
                            )
                            ch = sum(
                                table.rows[i].height
                                for i in range(row_index, row_index + cell.span_height)
                            )
                            frames.append(
                                (cell.text_frame, left, top, cw, ch, [row_index, column_index])
                            )
                        left += column.width
                    top += row.height
            for frame, bx, by, bw, bh, cell in frames:
                if not frame.text.strip():
                    continue
                # A shape's own rotation turns it about its centre, inside the group.
                turn = math.radians(shape.rotation) if cell is None else 0.0
                cs, sn = math.cos(turn), math.sin(turn)
                mx, my = bx + bw / 2, by + bh / 2
                corners = [
                    (
                        a * qx + c * qy + e,
                        b * qx + d * qy + f,
                    )
                    for qx, qy in (
                        (mx + (px - mx) * cs - (py - my) * sn, my + (px - mx) * sn + (py - my) * cs)
                        for px, py in [(bx, by), (bx + bw, by), (bx, by + bh), (bx + bw, by + bh)]
                    )
                ]
                rect = [
                    min(p[0] for p in corners),
                    min(p[1] for p in corners),
                    max(p[0] for p in corners),
                    max(p[1] for p in corners),
                ]
                # The unrotated frame in slide space: its side lengths and the angle of its top edge.
                width = math.dist(corners[0], corners[1]) / 12700
                height = math.dist(corners[0], corners[2]) / 12700
                angle = math.degrees(
                    math.atan2(corners[1][1] - corners[0][1], corners[1][0] - corners[0][0])
                )
                centre = ((rect[0] + rect[2]) / 25400, (rect[1] + rect[3]) / 25400)
                alignment = frame.paragraphs[0].alignment
                boxes.append(
                    {
                        "id": "shape-"
                        + "-".join(map(str, current_path))
                        + (f"-cell-{cell[0]}-{cell[1]}" if cell else ""),
                        "shape_id": shape.shape_id,
                        "shape_path": current_path,
                        "cell": cell,
                        "source_text": frame.text.replace("\v", "\n"),
                        "bbox": [v / 12700 for v in rect],
                        "frame": [centre[0] - width / 2, centre[1] - height / 2, width, height],
                        "angle": angle,
                        "rotation": shape.rotation,
                        "alignment": "center"
                        if alignment == 2
                        else "right"
                        if alignment == 3
                        else "left",
                        "vertical_alignment": "middle"
                        if frame.vertical_anchor == 3
                        else "bottom"
                        if frame.vertical_anchor == 4
                        else "top",
                        "margins": [
                            v / 12700
                            for v in [
                                frame.margin_left,
                                frame.margin_top,
                                frame.margin_right,
                                frame.margin_bottom,
                            ]
                        ],
                    }
                )

    visit(slide.shapes, (1, 0, 0, 1, 0, 0), [])
    return boxes


def match_box(trace: dict, text: str, boxes: list[dict]) -> dict | None:
    x0, y0, x1, y1 = trace["bbox"]
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    candidates = []
    for box in boxes:
        bx0, by0, bx1, by1 = box["bbox"]
        if not (bx0 - 3 <= cx <= bx1 + 3 and by0 - 3 <= cy <= by1 + 3):
            continue
        if normalize(text) not in normalize(box["source_text"]):
            continue
        candidates.append(((bx1 - bx0) * (by1 - by0), box))
    if candidates:
        return min(candidates, key=lambda item: item[0])[1]
    # Native PowerPoint sometimes lets a paragraph extend below its nominal box.
    # Allow that overflow only when the actual source text agrees.
    nearby = []
    for box in boxes:
        bx0, by0, bx1, by1 = box["bbox"]
        dx, dy = max(bx0 - cx, 0, cx - bx1), max(by0 - cy, 0, cy - by1)
        if dx > 3 or dy > trace["size"] * 2:
            continue
        native, source = normalize(text), normalize(box["source_text"])
        if native in source:
            nearby.append((dx + dy, box))
        elif dx == 0 and dy == 0 and len(native) >= 5:
            # PDF ligature decoding may differ from the PPTX Unicode source.
            # Retain both strings; do not use the decoded PDF as the translation source.
            strings = [source, *[normalize(word) for word in box["source_text"].split()]]
            score = max(SequenceMatcher(None, native, value).ratio() for value in strings)
            if score >= 0.8:
                nearby.append((100 * (1 - score), box))
    nearby.sort(key=lambda item: item[0])
    if not nearby or (len(nearby) > 1 and abs(nearby[0][0] - nearby[1][0]) < 0.1):
        return None
    return nearby[0][1]


def page_blocks(lines: list[TextLine], boxes: list[dict]) -> list[dict]:
    """Text blocks of one page: PowerPoint text frames where known, otherwise PDF paragraphs."""
    claimed: dict[str, list[TextLine]] = defaultdict(list)
    free = []
    for line in lines:
        trace = {"bbox": [line.x0, line.y0, line.x1, line.y1], "size": line.size}
        box = match_box(trace, line.text, boxes) if boxes else None
        if box:
            claimed[box["id"]].append(line)
        else:
            free.append(line)
    # A line the text could not place (a shared edge, a decoded ligature) still belongs
    # to a frame that already owns its neighbours, never to a second block.
    owners = [box for box in boxes if box["id"] in claimed]
    unplaced = []
    for line in free:
        cx, cy = (line.x0 + line.x1) / 2, (line.y0 + line.y1) / 2
        inside = [
            box
            for box in owners
            if box["bbox"][0] - 3 <= cx <= box["bbox"][2] + 3
            and box["bbox"][1] - 3 <= cy <= box["bbox"][3] + 3
        ]
        if inside:
            smallest = min(
                inside, key=lambda b: (b["bbox"][2] - b["bbox"][0]) * (b["bbox"][3] - b["bbox"][1])
            )
            claimed[smallest["id"]].append(line)
        else:
            unplaced.append(line)
    free = unplaced
    blocks = []
    for box in boxes:
        matched = claimed.get(box["id"])
        if matched:
            blocks.append(frame_block(box, matched))
    blocks.extend(pdf_block(block) for block in group_blocks(free))
    blocks.sort(key=lambda b: (round(b["cover"][1]), b["cover"][0]))
    return blocks


def cover_of(lines: list[TextLine]) -> list[float]:
    return [
        min(line.x0 for line in lines),
        min(line.y0 for line in lines),
        max(line.x1 for line in lines),
        max(line.y1 for line in lines),
    ]


def frame_block(box: dict, lines: list[TextLine]) -> dict:
    size = median(line.size for line in lines)
    tops = sorted(line.y0 for line in lines if line.rotation == 0)
    steps = [b - a for a, b in zip(tops, tops[1:], strict=False) if 0.5 * size < b - a < 3 * size]
    left, top, right, bottom = box["margins"]
    return {
        "source_id": box["id"],
        "text": box["source_text"],
        "frame": box["frame"],
        "angle": box["angle"],
        "padding": [top, right, bottom, left],
        "align": box["alignment"],
        "valign": box["vertical_alignment"],
        "size": size,
        "pitch": median(steps) if steps else size * 1.2,
        "color": next((line.color for line in lines if line.color), None),
        "cover": cover_of(lines),
    }


def pdf_block(block) -> dict:
    cx, cy = (block.x0 + block.x1) / 2, (block.y0 + block.y1) / 2
    width, height = block.x1 - block.x0, block.y1 - block.y0
    if block.rotation in (90, 270):
        width, height = height, width
    return {
        "source_id": None,
        "text": block.text,
        "frame": [cx - width / 2, cy - height / 2, width, height],
        "angle": block.rotation,
        "padding": [0, 0, 0, 0],
        "align": "left",
        "valign": "top",
        "size": block.size,
        "pitch": block.pitch,
        "color": block.color,
        "cover": cover_of(block.lines),
    }


def dominant_color(
    image: Image.Image, scale: tuple[float, float], cover: list[float]
) -> tuple[tuple[int, int, int], float]:
    """Most common colour under a text area and the share of pixels near it."""
    left, top = (
        max(0, math.floor(cover[0] * scale[0]) - 1),
        max(0, math.floor(cover[1] * scale[1]) - 1),
    )
    right = min(image.width, math.ceil(cover[2] * scale[0]) + 1)
    bottom = min(image.height, math.ceil(cover[3] * scale[1]) + 1)
    if right <= left or bottom <= top:
        return (255, 255, 255), 0.0
    crop = image.crop((left, top, right, bottom))
    if crop.width * crop.height > SAMPLE_PIXELS:
        shrink = math.sqrt(SAMPLE_PIXELS / (crop.width * crop.height))
        crop = crop.resize(
            (max(1, int(crop.width * shrink)), max(1, int(crop.height * shrink))),
            Image.Resampling.NEAREST,
        )
    counts: Counter[tuple[int, int, int]] = Counter()
    sums: dict[tuple[int, int, int], list[int]] = defaultdict(lambda: [0, 0, 0])
    data = crop.tobytes()
    pixels = list(zip(data[0::3], data[1::3], data[2::3], strict=True))
    for red, green, blue in pixels:
        key = (red >> 4, green >> 4, blue >> 4)
        counts[key] += 1
        total = sums[key]
        total[0] += red
        total[1] += green
        total[2] += blue
    key, count = counts.most_common(1)[0]
    red, green, blue = (round(v / count) for v in sums[key])
    return (red, green, blue), count / len(pixels)


def luminance(colour: tuple[int, int, int]) -> float:
    def channel(value: int) -> float:
        level = value / 255
        return level / 12.92 if level <= 0.03928 else ((level + 0.055) / 1.055) ** 2.4

    red, green, blue = (channel(v) for v in colour)
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def readable_colour(
    foreground: tuple[int, int, int] | None, background: tuple[int, int, int]
) -> tuple[int, int, int]:
    """The PDF's own text colour, unless it would not be legible on the sampled background."""
    shade = luminance(background)
    if foreground is not None:
        light, dark = sorted((luminance(foreground), shade), reverse=True)
        if (light + 0.05) / (dark + 0.05) >= 3:
            return foreground
    return (17, 17, 17) if shade > 0.4 else (255, 255, 255)


def percent(value: float, total: float) -> str:
    return f"{value / total * 100:.4f}"


def block_html(
    number: int,
    index: int,
    block: dict,
    page: tuple[float, float],
    image: Image.Image,
    scale: tuple[float, float],
) -> str:
    width, height = page
    x, y, w, h = block["frame"]
    top, right, bottom, left = block["padding"]
    size = max(block["size"], 1.0)
    background, share = dominant_color(image, scale, block["cover"])
    foreground = readable_colour(block["color"], background)
    fill = (
        "#%02x%02x%02x" % background
        if share >= UNIFORM_SHARE
        else "rgba(%d,%d,%d,.88)" % background
    )
    line_height = min(max(block["pitch"] / size, 1.0), 1.6)
    justify = {"middle": "center", "bottom": "flex-end"}.get(block["valign"], "flex-start")
    style = (
        f"left:{percent(x, width)}%;top:{percent(y, height)}%;"
        f"width:{percent(w, width)}%;height:{percent(h, height)}%;"
        f"font-size:{percent(size, width)}cqw;line-height:{line_height:.3f};"
        f"padding:{percent(top, width)}cqw {percent(right, width)}cqw "
        f"{percent(bottom, width)}cqw {percent(left, width)}cqw;"
        f"justify-content:{justify};text-align:{block['align']};"
        f"--fg:#{'%02x%02x%02x' % foreground};--bg:{fill}"
    )
    if abs(block["angle"]) > 0.05:
        style += f";transform:rotate({block['angle']:.3f}deg)"
    attributes = {
        "class": "pdf-text",
        "id": f"p{number}-text-{index}",
        "data-pw": f"{width:.3f}",
        "data-cw": f"{max(w - left - right, 1.0):.3f}",
        "data-ch": f"{max(h - top - bottom, 1.0):.3f}",
        "data-size": f"{size:.3f}",
        "data-lh": f"{line_height:.3f}",
        "style": style,
    }
    if block["source_id"]:
        attributes["data-source-id"] = block["source_id"]
    rendered = " ".join(f'{name}="{html.escape(value)}"' for name, value in attributes.items())
    return f"<div {rendered}>{html.escape(block['text'], quote=False)}</div>"


def page_html(
    number: int,
    page: tuple[float, float],
    blocks: list[dict],
    image: Image.Image,
    scale: tuple[float, float],
) -> str:
    width, height = page
    layer = "".join(
        block_html(number, index, block, page, image, scale)
        for index, block in enumerate(blocks, 1)
    )
    return (
        '<!doctype html><html lang="ja"><head><meta charset="utf-8">'
        f"<style>{PAGE_STYLE}</style></head><body>"
        f'<div class="page-canvas" data-number="{number}" data-width="{width:g}" data-height="{height:g}">'
        f'<img src="page-{number}.webp" alt="" draggable="false">{layer}</div></body></html>'
    )


def render_page_image(page) -> tuple[Image.Image, tuple[float, float]]:
    """The page as an RGB image and its pixels per point; the caller holds the PDFium lock."""
    width, height = page.get_size()
    scale = min(
        IMAGE_SCALE,
        IMAGE_MAX_SIDE / max(width, height),
        math.sqrt(IMAGE_MAX_PIXELS / (width * height)),
    )
    bitmap = page.render(scale=max(scale, 0.05))
    try:
        image = bitmap.to_pil().convert("RGB")
    finally:
        bitmap.close()
    return image, (image.width / width, image.height / height)


def build_editable_preview(source: Path, pdf: Path, folder: Path) -> str:
    pdf_relative = (
        "original.pdf"
        if pdf.resolve() == original_file(folder, ".pdf").resolve()
        else pdf.resolve().relative_to(folder.resolve()).as_posix()
    )
    deck = Presentation(str(source)) if source.suffix.lower() == ".pptx" else None
    if deck is not None and (deck.slide_width is None or deck.slide_height is None):
        raise ValueError("PowerPointのページ寸法がありません。")
    original_digest, pdf_digest = digest(source), digest(pdf)
    directory = folder / "editable-preview"
    info = checked_pdf(pdf)
    if deck is not None and info.pages != len(deck.slides):
        raise ValueError("PowerPointとPDFのページ数が一致しません。")
    if deck is not None and any(
        abs(width - float(deck.slide_width or 0) / 12700) > 0.1
        or abs(height - float(deck.slide_height or 0) / 12700) > 0.1
        for width, height in info.sizes
    ):
        raise ValueError("PowerPointとPDFのページ寸法が一致しません。")
    # Only one page's image and text are alive at a time. Publish the complete
    # directory only after every page succeeds.
    with TemporaryDirectory(prefix=".editable-preview-", suffix=".tmp", dir=folder) as work:
        staging = Path(work) / "editable-preview"
        staging.mkdir()
        assets = ["Preview.html"]
        blocks_count = mapped_count = 0
        with pypdfium2_lock:
            document = open_pdf(pdf)
        try:
            with (staging / "Preview.html").open("w") as preview:
                preview.write(
                    '<!doctype html><html lang="ja"><head><meta charset="utf-8">'
                    f"<title>{html.escape(source.name)}</title><style>body{{margin:0;background:#e8edf2}}.slide{{margin:12px auto;background:white}}iframe{{width:100%;height:100%;border:0}}</style></head><body>"
                )
                for number, size in enumerate(info.sizes, 1):
                    with pypdfium2_lock:
                        page = document[number - 1]
                        try:
                            image, scale = render_page_image(page)
                            lines = read_page(page)
                        finally:
                            page.close()
                    boxes = source_textboxes(deck.slides[number - 1]) if deck is not None else []
                    blocks = page_blocks(lines, boxes)
                    image.save(
                        staging / f"page-{number}.webp", "WEBP", quality=IMAGE_QUALITY, method=4
                    )
                    (staging / f"page-{number}.html").write_text(
                        page_html(number, size, blocks, image, scale)
                    )
                    image.close()
                    preview.write(
                        f'<div class="slide" data-number="{number}" style="width:{size[0]:g}px;height:{size[1]:g}px"><iframe loading="lazy" src="page-{number}.html" title="ページ {number}" sandbox></iframe></div>'
                    )
                    assets.extend((f"page-{number}.html", f"page-{number}.webp"))
                    blocks_count += len(blocks)
                    mapped_count += sum(bool(b["source_id"]) for b in blocks)
                preview.write("</body></html>")
        finally:
            with pypdfium2_lock:
                document.close()
        if digest(source) != original_digest or digest(pdf) != pdf_digest:
            raise ValueError("生成中に原本またはPDFが変更されました。")
        manifest = {
            "version": RENDERER_VERSION,
            "renderer": "PDFium page image / text layer",
            "pdf": pdf_relative,
            "source_sha256": original_digest,
            "pdf_sha256": pdf_digest,
            "pages": info.pages,
            "assets": assets,
            "text_blocks": blocks_count,
            "mapped_blocks": mapped_count,
        }
        (staging / "manifest.json").write_text(json.dumps(manifest, indent=2))
        previous = folder / f".editable-preview-old-{uuid4().hex}.tmp"
        if directory.exists():
            directory.rename(previous)
        try:
            staging.rename(directory)
        except OSError:
            if previous.exists():
                previous.rename(directory)
            raise
        finally:
            if previous.exists() and directory.exists():
                shutil.rmtree(previous)
    return "editable-preview/Preview.html"


def ensure_pdf_preview(source: Path, folder: Path) -> str:
    with PDF_RENDER_LOCK:
        return editable_preview(source, folder) or build_editable_preview(source, source, folder)


def refresh_saved_preview(source: Path, folder: Path) -> str | None:
    """Rebuild a PowerPoint preview from an older renderer, from its saved PDF.

    Nothing is exported: only the PDF a verified earlier export left in the folder
    is used, so the host needs no Office. Returns None when that PDF is missing or
    no longer matches; the caller then treats the preview as unavailable.
    """
    from docling_desk.preview.powerpoint_export import saved_export

    with PDF_RENDER_LOCK:
        if current := editable_preview(source, folder):
            return current
        try:
            pdf = saved_export(source, folder)
        except (OSError, ValueError, RuntimeError):
            return None
        return build_editable_preview(source, pdf, folder)
