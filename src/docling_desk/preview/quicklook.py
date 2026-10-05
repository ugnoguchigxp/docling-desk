"""Repair known Quick Look PPTX HTML defects without modifying the source deck."""

from __future__ import annotations

import json
import logging
import math
import re
from pathlib import Path
from statistics import median
from urllib.parse import unquote, urlsplit

import pypdfium2 as pdfium
from docling.utils.locks import pypdfium2_lock
from lxml import html
from pptx import Presentation

LOGGER = logging.getLogger(__name__)
EMU_PER_POINT = 12700


def declarations(value: str) -> dict[str, str]:
    return {
        key.strip().lower(): item.strip()
        for entry in value.split(";")
        if ":" in entry
        for key, item in [entry.split(":", 1)]
    }


def number(value: str | None) -> float | None:
    match = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)(?:px)?\s*", value or "")
    return float(match[1]) if match else None


def inline(element, **properties: str) -> None:
    # Quick Look uses repeated class selectors; inline !important wins reliably.
    style = declarations(element.get("style", ""))
    style.update(
        {key.replace("_", "-"): value + " !important" for key, value in properties.items()}
    )
    element.set("style", ";".join(f"{key}:{value}" for key, value in style.items()))


def normalize(text: str) -> str:
    return " ".join(text.split())


def rasterize_pdf(source: Path, target: Path) -> None:
    """Keep transparency and limit memory; share PDFium's lock with Docling."""
    with pypdfium2_lock, pdfium.PdfDocument(source) as document:
        if len(document) != 1:
            raise ValueError("Quick Look diagram must have exactly one PDF page")
        page = document[0]
        try:
            width, height = page.get_size()
            scale = min(3.0, math.sqrt(12_000_000 / (width * height)), 8192 / max(width, height))
            bitmap = page.render(scale=scale, fill_color=(0, 0, 0, 0))
            try:
                bitmap.to_pil().save(target, format="PNG")
            finally:
                bitmap.close()
        finally:
            page.close()


def prune_repeated_triangles(slides: list, deck) -> list[dict]:
    """Reject a reused triangle on pages that have no triangle in their source.

    Quick Look wraps these vectors in a 20px transparent inset and sometimes
    repeats the same PDF and coordinates on unrelated slides. Only discard an
    instance after finding its source owner; unknown vectors stay untouched.
    """
    owners: dict[tuple, list[int]] = {}
    occurrences = []
    has_triangle = []
    for index, (slide, native) in enumerate(zip(slides, deck.slides, strict=True), 1):
        scopes = (
            native.shapes,
            native.slide_layout.shapes,
            native.slide_layout.slide_master.shapes,
        )
        has_triangle.append(
            any(
                shape._element.xpath('.//a:prstGeom[@prst="triangle"]')
                for shapes in scopes
                for shape in shapes
            )
        )
        triangles = [
            shape
            for shapes in scopes
            for shape in shapes
            if not shape.rotation and shape._element.xpath('./p:spPr/a:prstGeom[@prst="triangle"]')
        ]
        for image in slide.xpath("./img[@src]"):
            if urlsplit(image.get("src")).path.lower().endswith(".pdf"):
                properties = declarations(image.get("style", ""))
                box = tuple(
                    number(properties.get(key)) for key in ("left", "top", "width", "height")
                )
                if any(value is None for value in box):
                    continue
                key = (image.get("src"), box)
                occurrences.append((index, image, key))
                for shape in triangles:
                    expected = (
                        shape.left / EMU_PER_POINT - 20,
                        shape.top / EMU_PER_POINT - 20,
                        shape.width / EMU_PER_POINT + 40,
                        shape.height / EMU_PER_POINT + 40,
                    )
                    if all(abs(a - b) < 1.5 for a, b in zip(box, expected, strict=True)):
                        owners.setdefault(key, []).append(index)
                        break
    removed = []
    for index, image, key in occurrences:
        if key in owners and not has_triangle[index - 1]:
            removed.append({"slide": index, "asset": key[0], "source_slides": owners[key]})
            image.drop_tree()
    return removed


def corrected_preview(source: Path, preview: Path) -> Path:
    """Produce a separate derived preview, retaining untouched Quick Look output."""
    tree = html.fromstring(preview.read_text(encoding="utf-8"))
    css = "\n".join(tree.xpath("//style/text()"))
    classes = {}
    # Generated Quick Look styles contain simple class rules, without nested rules.
    for selectors, body in re.findall(r"([^{}]+)\{([^{}]*)\}", css):
        for selector in selectors.split(","):
            if re.fullmatch(r"(?:\.[\w-]+)+\s*", selector.strip()):
                for name in re.findall(r"\.([\w-]+)", selector):
                    classes.setdefault(name, {}).update(declarations(body))

    def computed(element):
        values = {}
        for name in element.get("class", "").split():
            values.update(classes.get(name, {}))
        values.update(declarations(element.get("style", "")))
        return values

    slides = tree.xpath('//div[contains(concat(" ",normalize-space(@class)," ")," slide ")]')
    deck = Presentation(str(source))
    if len(slides) != len(deck.slides):
        return preview  # Never transfer formatting between different pages.
    report = {
        "text_frames": 0,
        "source_matched_frames": 0,
        "pdf_assets": 0,
        "errors": [],
        "removed_cross_slide_shapes": prune_repeated_triangles(slides, deck),
    }
    for slide, native in zip(slides, deck.slides, strict=True):
        for frame in slide.xpath("./div[@style]"):
            paragraphs = frame.xpath(".//p")
            if not paragraphs:
                continue
            properties = computed(frame)
            box = [number(properties.get(key)) for key in ("left", "top", "width", "height")]
            if any(value is None for value in box):
                continue
            box = [value for value in box if value is not None]
            candidates = []
            for shape in native.shapes:
                if not shape.has_text_frame or shape.rotation:
                    continue
                coordinates = [
                    v / EMU_PER_POINT for v in (shape.left, shape.top, shape.width, shape.height)
                ]
                if all(abs(a - b) < 1.5 for a, b in zip(box, coordinates, strict=True)):
                    candidates.append(shape)
            shape = candidates[0] if len(candidates) == 1 else None
            native_paragraphs = list(shape.text_frame.paragraphs) if shape is not None else []
            # Paragraph count AND text must agree before using source formatting.
            matched = len(native_paragraphs) == len(paragraphs) and all(
                normalize(p.text) == normalize("".join(q.itertext()).lstrip("•·▪◦–-\u00a0 "))
                or normalize(p.text) == normalize("".join(q.itertext()))
                for p, q in zip(native_paragraphs, paragraphs, strict=False)
            )
            inner = frame.find("div")
            if inner is not None and computed(inner).get("display") == "table":
                values = computed(inner)
                left, top, right, bottom = [
                    number(values.get(f"margin-{side}")) or 0
                    for side in ("left", "top", "right", "bottom")
                ]
                if matched and shape is not None:
                    text_frame = shape.text_frame
                    left, top, right, bottom = [
                        v / EMU_PER_POINT
                        for v in (
                            text_frame.margin_left,
                            text_frame.margin_top,
                            text_frame.margin_right,
                            text_frame.margin_bottom,
                        )
                    ]
                inline(
                    inner,
                    width=f"{max(0, box[2] - left - right):.3f}px",
                    height=f"{max(0, box[3] - top - bottom):.3f}px",
                    margin=f"{top:.3f}px {right:.3f}px {bottom:.3f}px {left:.3f}px",
                    table_layout="fixed",
                )
                report["text_frames"] += 1
            nonempty = [i for i, p in enumerate(paragraphs) if normalize("".join(p.itertext()))]
            for index, paragraph in enumerate(paragraphs):
                props = computed(paragraph)
                spans = paragraph.xpath(".//span")
                bullet = (
                    spans[0]
                    if spans
                    and normalize("".join(spans[0].itertext())) in {"•", "·", "▪", "◦", "–", "-"}
                    else None
                )
                sizes = [
                    number(computed(span).get("font-size")) for span in spans if span is not bullet
                ]
                sizes = [size for size in sizes if size is not None and size > 0]
                font_size = median(sizes) if sizes else number(props.get("font-size"))
                if font_size:
                    family = next(
                        (
                            computed(span).get("font-family")
                            for span in spans
                            if span is not bullet and computed(span).get("font-family")
                        ),
                        props.get("font-family", "sans-serif"),
                    )
                    inline(
                        paragraph,
                        font_size=f"{font_size:g}px",
                        font_family=f"{family},Arial,sans-serif",
                    )
                    # Match the run font to avoid a 24px paragraph strut around 14px text.
                    for span in spans:
                        family = computed(span).get("font-family")
                        if family:
                            inline(span, font_family=f"{family},Arial,sans-serif")
                if matched:
                    original = native_paragraphs[index]
                    spacing = original.line_spacing
                    if isinstance(spacing, float):
                        inline(paragraph, line_height=f"{spacing * 1.2:g}")
                    elif spacing is not None:
                        inline(paragraph, line_height=f"{spacing.pt:g}px")
                    else:
                        inline(paragraph, line_height="1.2")
                    before = original.space_before.pt if original.space_before else 0
                    after = original.space_after.pt if original.space_after else 0
                    inline(
                        paragraph,
                        padding_top=f"{before if index else 0:g}px",
                        padding_bottom=f"{after:g}px",
                    )
                    # PowerPoint's centered boxes ignore empty boundary paragraphs.
                    # Retain their nodes/text, but avoid adding browser-only blank lines.
                    if nonempty and (index < nonempty[0] or index > nonempty[-1]):
                        inline(
                            paragraph,
                            font_size="0px",
                            line_height="0",
                            padding_top="0px",
                            padding_bottom="0px",
                        )
                if bullet is not None:
                    indent = -(number(props.get("text-indent")) or 0)
                    if matched:
                        ppr = native_paragraphs[index]._p.pPr
                        if ppr is not None and ppr.get("indent"):
                            indent = -int(ppr.get("indent")) / EMU_PER_POINT
                    if indent > 0:
                        inline(paragraph, margin_left=f"{indent:g}px", text_indent=f"{-indent:g}px")
                        inline(
                            bullet,
                            margin_left="0px",
                            padding_right="0px",
                            display="inline-block",
                            width=f"{indent:g}px",
                            text_indent="0px",
                            text_align="left",
                        )
            if matched:
                report["source_matched_frames"] += 1

    converted: dict[str, str | None] = {}
    for image in tree.xpath("//img[@src]"):
        url = urlsplit(image.get("src"))
        if url.scheme or url.netloc or url.query or url.fragment:
            continue
        name = unquote(url.path)
        path = (preview.parent / name).resolve()
        if path.suffix.lower() != ".pdf" or not path.is_relative_to(preview.parent.resolve()):
            continue
        if name not in converted:
            target = path.with_name(path.stem + ".quicklook.png")
            try:
                rasterize_pdf(path, target)
                converted[name] = target.relative_to(preview.parent.resolve()).as_posix()
                report["pdf_assets"] += 1
            except Exception as exc:
                LOGGER.warning("Unable to render Quick Look diagram %s: %s", path.name, exc)
                converted[name] = None
                report["errors"].append({"asset": name, "error": str(exc)})
        if converted[name]:
            image.set("src", converted[name])
    # Keep the original quirks-mode rendering: Quick Look uses unitless lengths.
    target = preview.with_name("Preview-corrected.html")
    target.write_text(html.tostring(tree, encoding="unicode"), encoding="utf-8")
    preview.with_name("preview-corrections.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return target
