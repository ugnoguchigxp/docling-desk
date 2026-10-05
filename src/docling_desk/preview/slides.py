"""Position Docling content on its source slide, without claiming pixel fidelity."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from statistics import median
from typing import Literal

from docling_core.types.doc import DoclingDocument, PictureItem, TableItem, TextItem
from lxml import html
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pydantic import BaseModel, Field

EMU_PER_POINT = 12700


class SlideCell(BaseModel):
    text: str
    row: int
    col: int
    row_span: int
    col_span: int
    header: bool


class SlideBlock(BaseModel):
    kind: Literal["text", "table", "picture"]
    refs: list[str] = Field(default_factory=list)
    texts: list[str] = Field(default_factory=list)
    x: float
    y: float
    width: float
    height: float
    font_size: float = 20
    cells: list[SlideCell] = Field(default_factory=list)
    rows: int = 0
    columns: int = 0
    image: str | None = None


class SlidePage(BaseModel):
    number: int
    width: float
    height: float
    preview: str | None = None
    blocks: list[SlideBlock] = Field(default_factory=list)
    unplaced: list[dict[str, str]] = Field(default_factory=list)


class SlideDocument(BaseModel):
    slides: list[SlidePage]
    notice: str = (
        "抽出配置はDoclingの座標を使った再構成です。原本＋抽出位置で矢印・背景・装飾も確認できます。"
        "フォント・重なり・回転・グループ図形の再現と、図の意味関係の自動理解は保証しません。"
    )


def shape_font_sizes(source: Path) -> dict[tuple[int, int, int, int, int], float]:
    """Reuse explicit PPTX text sizes; inheritance and themes remain approximate."""
    presentation = Presentation(str(source))
    sizes: dict[tuple[int, int, int, int, int], float] = {}

    def visit(shape: BaseShape, page: int) -> None:
        if isinstance(shape, GroupShape):
            for child in shape.shapes:
                visit(child, page)
        values: list[float] = []
        if isinstance(shape, Shape) and shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                if paragraph.font.size is not None:
                    values.append(paragraph.font.size.pt)
                for run in paragraph.runs:
                    if run.font.size is not None:
                        values.append(run.font.size.pt)
        if values:
            sizes[(page, shape.left, shape.top, shape.width, shape.height)] = median(values)

    for number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            visit(shape, number)
    return sizes


def split_quicklook(preview: Path, folder: Path, page_count: int) -> dict[int, str]:
    """Keep source rendering in script-blocked local frames, using shared existing assets."""
    tree = html.fromstring(preview.read_text(encoding="utf-8"))
    slides = tree.xpath('//div[contains(concat(" ",normalize-space(@class)," ")," slide ")]')
    if len(slides) != page_count:
        return {}  # A source slide must never be silently paired with the wrong page.
    styles = "\n".join(style.text or "" for style in tree.xpath("//style"))
    styles += "\nhtml,body{margin:0!important;padding:0!important;zoom:1!important;background:white!important;overflow:hidden}.slide{margin:0!important;box-shadow:none!important}"
    (preview.parent / "slide-layout.css").write_text(styles, encoding="utf-8")
    paths: dict[int, str] = {}
    for number, slide in enumerate(slides, start=1):
        for unsafe in slide.xpath(".//script|.//iframe|.//object|.//embed"):
            unsafe.drop_tree()
        for element in slide.iter():
            for attribute in list(element.attrib):
                if attribute.lower().startswith("on"):
                    del element.attrib[attribute]
        target = preview.parent / f"slide-layout-{number}.html"
        # Quick Look's unitless CSS requires its existing quirks-mode rendering.
        target.write_text(
            '<html><head><meta charset="utf-8"><link rel="stylesheet" href="slide-layout.css"></head><body>'
            + html.tostring(slide, encoding="unicode")
            + "</body></html>",
            encoding="utf-8",
        )
        paths[number] = target.relative_to(folder).as_posix()
    return paths


def export_slide_layout(
    doc: DoclingDocument, source: Path, folder: Path, preview: str | None
) -> None:
    fonts = shape_font_sizes(source)
    previews = split_quicklook(folder / preview, folder, len(doc.pages)) if preview else {}
    pages = {
        number: SlidePage(
            number=number,
            width=page.size.width / EMU_PER_POINT,
            height=page.size.height / EMU_PER_POINT,
            preview=previews.get(number),
        )
        for number, page in sorted(doc.pages.items())
    }
    groups: dict[tuple[int, str, float, float, float, float], SlideBlock] = {}
    unknown: dict[int, list[dict[str, str]]] = defaultdict(list)
    for item, _ in doc.iterate_items():
        if not isinstance(item, (TextItem, TableItem, PictureItem)):
            continue
        parent = item.parent.resolve(doc) if item.parent else None
        if isinstance(parent, (TableItem, PictureItem)):
            continue  # Captions/annotations are already represented with their parent.
        text = item.text if isinstance(item, TextItem) else item.label.value
        if not item.prov:
            unknown[0].append({"ref": item.self_ref, "text": text})
            continue
        prov = item.prov[0]
        if prov.page_no not in pages:
            unknown[0].append({"ref": item.self_ref, "text": text})
            continue
        size = doc.pages[prov.page_no].size
        box = prov.bbox.to_top_left_origin(page_height=size.height)
        x, y, width, height = box.l, box.t, box.r - box.l, box.b - box.t
        if width <= 0 or height <= 0:
            unknown[prov.page_no].append({"ref": item.self_ref, "text": text})
            continue
        kind = (
            "text"
            if isinstance(item, TextItem)
            else "table"
            if isinstance(item, TableItem)
            else "picture"
        )
        key = (prov.page_no, kind, x, y, width, height)
        block = groups.get(key)
        if block is None:
            block = SlideBlock(
                kind=kind,
                x=x / EMU_PER_POINT,
                y=y / EMU_PER_POINT,
                width=width / EMU_PER_POINT,
                height=height / EMU_PER_POINT,
                font_size=fonts.get(
                    (prov.page_no, round(x), round(y), round(width), round(height)),
                    32 if item.label.value in {"title", "section_header"} else 20,
                ),
            )
            groups[key] = block
            pages[prov.page_no].blocks.append(block)
        block.refs.append(item.self_ref)
        if isinstance(item, TextItem):
            block.texts.append(item.text)
        elif isinstance(item, TableItem):
            block.rows, block.columns = item.data.num_rows, item.data.num_cols
            block.font_size = 16
            block.cells = [
                SlideCell(
                    text=cell.text,
                    row=cell.start_row_offset_idx,
                    col=cell.start_col_offset_idx,
                    row_span=cell.row_span,
                    col_span=cell.col_span,
                    header=cell.column_header,
                )
                for cell in item.data.table_cells
            ]
        elif item.image is not None:
            uri = str(item.image.uri)
            if uri.startswith(
                ("data:image/png;", "data:image/jpeg;", "data:image/webp;", "data:image/gif;")
            ):
                block.image = uri
    for number, records in unknown.items():
        if number in pages:
            pages[number].unplaced.extend(records)
        elif pages:
            pages[min(pages)].unplaced.extend(records)
    (folder / "slides.json").write_text(
        SlideDocument(slides=list(pages.values())).model_dump_json(indent=2), encoding="utf-8"
    )
